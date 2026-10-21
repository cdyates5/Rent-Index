#!/usr/bin/env python3
"""
rebuild.py  -  Acheron All-Tenant Rent Index (A-ATRI)

A synthetic replacement for the BLS research series R-CPI-ATR ("all tenant regressed
rent"), whose publication BLS paused in April 2026 after the October-2025 Housing
Survey went uncollected. Regenerates the dashboard from an empty directory.

Method (no BLS microdata needed):
  1. Invert the official CPI rent estimator. CPI rent's monthly change is the sixth root
     of each panel's 6-month rent change, so 6*dln(CPI_t) = sum_k kappa_k g_{t-k}, where
     g is monthly growth of the all-tenant rent level dated to when rents were set (the
     BLS ATR re-dating rule) and kappa = p * box6 (p = distribution of re-dating offsets).
     g is recovered by Bayesian deconvolution (random-walk smoothness prior, lambda chosen
     by marginal likelihood). The Oct-2025 gap and the Apr-2026 12-month catch-up are
     modelled exactly as BLS computed them.
  2. Formula wedge alpha: BLS ATR is a log repeat-rent regression (geometric); CPI is a ratio
     of arithmetic means. alpha is pinned by the published 1999Q4 -> 2024Q4 level ratio.
     The re-dating weight w is chosen on published YoY anchors.
  3. Lease-roll layer: new-lease rents (Zillow ZORI) lead the all-tenant reset rate; direct
     h-step regressions supply an AR(1)-correlated prior for the ragged edge and forecast.
  4. Real-time backtest sets empirical forecast-band scaling.

Data (live first, embedded snapshot fallback):
  CPI rent NSA (CUUR0000SEHA): BLS API v2 (BLS_API_KEY optional) -> FRED API (FRED_API_KEY)
    -> DBnomics -> embedded snapshot.  Zillow ZORI national: Zillow CSV -> snapshot.
  BLS R-CPI-ATR/NTR xlsx: local file ./r-cpi-ntr-and-r-cpi-atr.xlsx (or ATR_XLSX env) or a
    download attempt. If found and parseable, the full published series replaces the anchors
    for calibration and validation. Otherwise published anchors (CEIC / press) are used.

Usage:
  python3 rebuild.py                                   local build ($ATRI_OUT, /mnt/user-data/outputs or ./output)
  python3 rebuild.py --out site --html-name index.html --strict --archive     CI build (see README)

  --strict   refuse to write anything unless CPI came from a live API (BLS or FRED), ZORI is live,
             and the result passes the sanity gate. Exit code 2 on failure.
  --archive  append this build's headline to state/headline_log.csv and save the quarterly table
             to state/vintages/ (real-time vintages of the synthetic index).

Environment: BLS_API_KEY, FRED_API_KEY (optional, one strongly recommended), BLS_CONTACT_EMAIL
(used in the User-Agent for www.bls.gov, which blocks anonymous clients), ATR_XLSX (path to the
BLS research workbook, optional; data/r-cpi-ntr-and-r-cpi-atr.xlsx is picked up automatically).
"""
import os, re, io, sys, csv, json, glob, math, argparse, hashlib, datetime as dt, urllib.request, urllib.error
import numpy as np
import pandas as pd

UA = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.environ.get("ATRI_OUT") or ("/mnt/user-data/outputs" if os.path.isdir("/mnt/user-data/outputs") else "./output")
HTML_NAME = "acheron_atri.html"
LIVE_CPI_SOURCES = ("BLS API v2", "FRED API")
CSV_NAME = "acheron_atri_quarterly.csv"
LOG = []


def log(msg):
    LOG.append(msg)
    print(msg, flush=True)


# ----------------------------------------------------------------------------------------
# Embedded snapshots (retrieved 2026-10-19). Live sources override these where available.
# ----------------------------------------------------------------------------------------
CPI_SNAP_START = "1997-01"   # CUUR0000SEHA NSA; 2025-10 uncollected (null). 2026-01..04 derived
CPI_SNAP = [164.400,164.800,165.100,165.500,165.900,166.400,166.800,167.300,167.800,168.200,168.700,169.100,169.500,169.900,170.300,170.700,171.100,171.700,172.200,172.800,173.400,173.900,174.500,174.900,175.300,175.600,176.000,176.400,176.700,177.100,177.500,177.900,178.400,178.800,179.800,180.300,181.100,181.500,182.000,182.300,182.700,183.200,183.900,184.600,185.300,186.100,186.800,187.600,188.200,188.900,189.600,190.200,191.000,191.600,192.300,193.100,193.900,194.700,195.500,196.400,197.000,197.700,198.200,198.500,198.800,199.300,199.800,200.200,200.700,201.300,202.000,202.500,203.300,203.700,204.100,204.500,204.900,205.100,205.600,206.100,206.600,206.900,207.500,207.900,208.300,208.800,209.200,209.700,210.200,210.700,211.200,211.900,212.400,212.800,213.200,213.900,214.500,215.000,215.500,216.000,216.400,216.800,217.500,218.000,218.600,219.300,220.000,220.500,220.900,221.600,222.300,222.900,223.600,224.400,225.200,226.200,227.100,228.000,228.900,230.000,230.806,231.739,232.495,232.980,233.549,234.071,234.732,235.311,236.058,237.135,238.169,239.102,239.850,240.325,240.874,241.474,241.803,242.640,243.367,244.181,244.926,245.855,246.681,247.278,247.974,248.305,248.639,248.899,249.069,249.092,248.994,249.029,248.965,248.888,248.886,248.999,249.144,249.017,249.089,249.012,248.925,248.999,249.126,249.024,249.368,249.618,250.317,250.986,251.555,251.829,252.145,252.221,252.393,252.592,253.085,254.003,254.628,255.651,256.367,257.189,257.714,258.184,258.569,258.922,259.231,259.407,260.107,260.677,261.421,262.707,263.365,264.098,264.700,265.256,265.821,265.984,266.559,266.905,267.482,268.505,269.137,269.960,270.698,271.688,272.317,272.733,273.486,274.100,274.710,275.321,276.248,277.048,277.998,278.985,280.123,280.874,281.572,282.389,283.130,283.598,284.245,285.031,286.090,287.068,288.306,289.428,290.322,291.204,292.004,292.777,293.489,294.175,295.036,295.902,296.862,297.916,298.962,300.400,301.587,302.735,303.467,304.211,304.868,305.477,306.379,307.314,308.173,309.479,310.268,311.501,312.670,313.904,314.788,315.277,315.883,316.763,317.490,318.318,319.351,320.651,321.533,322.628,323.968,324.815,325.597,326.351,327.513,328.678,329.333,330.648,331.605,332.638,333.834,334.680,335.819,336.789,337.825,338.616,339.519,340.135,340.811,341.294,341.950,342.444,342.910,343.615,344.039,344.455,344.758,345.242,345.717,346.267,347.016,347.833,348.469,349.710,351.255,352.892,354.526,355.931,357.737,359.627,361.083,362.951,365.116,367.927,370.448,373.283,376.569,379.436,382.562,385.649,388.372,391.141,392.910,394.898,396.726,398.593,400.210,402.247,404.487,406.683,408.838,410.606,412.019,413.695,415.219,416.386,417.772,418.820,420.577,422.223,423.821,425.381,426.651,428.151,429.506,430.603,431.798,432.956,433.698,434.594,435.489,436.981,438.212,None,439.275,440.667,441.716,442.158,442.862,445.029,446.380,446.945,447.963,448.997,450.402]    # from FRED SA x 2025 seasonal factors (validated to <0.002 on May-Sep)
ZORI_SNAP_START = "2015-01"  # Zillow ZORI, all homes, smoothed, United States
ZORI_SNAP = [1130.39,1136.69,1145.21,1153.92,1162.44,1169.86,1175.36,1178.85,1180.68,1180.96,1181.25,1181.64,1184.58,1189.18,1194.97,1202.84,1210.96,1218.96,1223.16,1225.03,1224.74,1223.40,1222.20,1221.58,1225.98,1233.46,1244.70,1254.17,1263.20,1270.03,1274.59,1275.97,1276.16,1274.51,1274.37,1275.10,1279.75,1286.88,1296.16,1305.40,1313.27,1320.43,1325.44,1328.24,1327.27,1326.58,1327.22,1329.35,1333.13,1338.87,1348.77,1358.71,1368.28,1375.83,1381.11,1383.68,1383.10,1381.44,1380.33,1381.49,1386.40,1394.13,1402.48,1404.16,1404.47,1405.68,1410.91,1412.22,1409.64,1406.60,1404.93,1407.00,1415.50,1427.86,1443.11,1460.70,1484.34,1512.21,1541.29,1568.34,1590.17,1603.77,1612.82,1621.86,1632.27,1646.57,1661.31,1680.66,1700.33,1720.58,1737.23,1746.62,1748.48,1743.86,1738.03,1735.46,1737.45,1744.08,1756.55,1767.10,1778.47,1784.78,1792.77,1796.81,1796.89,1793.32,1788.43,1786.26,1789.36,1796.84,1806.94,1817.16,1827.54,1836.57,1843.21,1845.89,1845.15,1841.44,1838.27,1836.49,1840.92,1848.85,1859.84,1868.56,1875.59,1880.41,1882.80,1882.92,1881.17,1878.19,1874.71,1872.61,1875.56,1883.06,1893.65,1903.90,1913.13,1920.64,1926.02,1929.79,1931.94]

# BLS R-CPI-NTR (current methodology, 2025Q3 vintage released 2026-01-29), via YCharts
NTR_SNAP = {"2013Q2": 138.02, "2013Q3": 138.68, "2013Q4": 138.90, "2014Q1": 140.43, "2014Q2": 141.54,
            "2014Q3": 143.94, "2014Q4": 143.62, "2015Q1": 145.49, "2015Q2": 147.12, "2015Q3": 150.59,
            "2015Q4": 149.95, "2016Q1": 149.99, "2016Q2": 153.11, "2016Q3": 154.35, "2016Q4": 153.38,
            "2017Q1": 154.84, "2017Q2": 158.48, "2017Q3": 158.52, "2017Q4": 159.12, "2018Q1": 159.00,
            "2018Q2": 162.35, "2018Q3": 164.30, "2018Q4": 162.61, "2019Q1": 164.98, "2019Q2": 167.19,
            "2019Q3": 170.43, "2019Q4": 169.22, "2020Q1": 170.80, "2020Q2": 170.08, "2020Q3": 171.08,
            "2020Q4": 172.15, "2021Q1": 173.00, "2021Q2": 177.45, "2021Q3": 183.66, "2021Q4": 192.15,
            "2022Q1": 194.61, "2022Q2": 200.83, "2022Q3": 203.27, "2022Q4": 202.15, "2023Q1": 203.57,
            "2023Q2": 207.81, "2023Q3": 209.40, "2023Q4": 206.99, "2024Q1": 209.02, "2024Q2": 212.80,
            "2024Q3": 212.92, "2024Q4": 209.54, "2025Q1": 212.27, "2025Q2": 208.70, "2025Q3": 215.60}

# Published R-CPI-ATR anchors. CEIC = BLS via CEIC, Dec-2024 vintage (pre-Apr-2025 method).
ANCHORS = [
    dict(q="1999Q4", kind="level", value=102.275, src="CEIC (BLS), Dec-2024 vintage", role="calibration"),
    dict(q="2010Q1", kind="yoy", value=-1.098, src="CEIC: series record low", role="w selection"),
    dict(q="2022Q4", kind="yoy", value=7.529, src="CEIC: series record high", role="w selection"),
    dict(q="2023Q4", kind="yoy", value=5.4, src="N. Timiraos (WSJ), Apr-2024 first release", role="check"),
    dict(q="2024Q1", kind="yoy", value=5.4, src="N. Timiraos (WSJ), Apr-2024 first release", role="check"),
    dict(q="2024Q3", kind="yoy", value=3.858, src="CEIC, Dec-2024 vintage", role="w selection"),
    dict(q="2024Q3", kind="level", value=200.459, src="CEIC, Dec-2024 vintage", role="check"),
    dict(q="2024Q4", kind="yoy", value=3.182, src="CEIC, Dec-2024 first release", role="w selection"),
    dict(q="2024Q4", kind="level", value=201.766, src="CEIC, Dec-2024 first release", role="calibration"),
]
ANCHOR_MEDIAN = 2.750          # CEIC: median YoY, 1999Q4-2024Q4 sample
ANCHOR_PEAK, ANCHOR_TROUGH = "2022Q4", "2010Q1"
LAST_BLS_PRINT = "2025Q3"      # last quarter BLS published (Jan-2026 release); paused Apr-2026

# MLR (2026) "interpolate" counterfactual: published Oct-2025 rent level 438.212 vs 439.361
CPI_GAP = math.log(439.361 / 438.212)     # published levels Oct-2025..Mar-2026 understated by this
GAP_MONTHS = ("2025-10", "2026-03")

W_GRID = np.round(np.arange(0.0, 1.0001, 0.05), 2)
PRIOR_FROM_LAG = 4             # lease-roll prior applies to the last 4 CPI months + forecast
H_FCST_M = 9                   # forecast horizon beyond last CPI month (months)


# ----------------------------------------------------------------------------------------
# Fetching
# ----------------------------------------------------------------------------------------
def http_get(url, data=None, headers=None, timeout=45):
    ua = UA
    if "bls.gov" in url and os.environ.get("BLS_CONTACT_EMAIL"):
        ua = f"Acheron-ATRI-rebuild/1.0 ({os.environ['BLS_CONTACT_EMAIL']})"   # BLS asks for a contact UA
    h = {"User-Agent": ua}
    h.update(headers or {})
    req = urllib.request.Request(url, data=data, headers=h)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def snap_series(start, vals):
    idx = pd.period_range(start, periods=len(vals), freq="M")
    return pd.Series([np.nan if v is None else float(v) for v in vals], index=idx)


def fetch_cpi():
    s = snap_series(CPI_SNAP_START, CPI_SNAP)
    src = ["embedded snapshot (2026-10-19)"]
    live = {}
    key = os.environ.get("BLS_API_KEY")
    try:
        span = 20 if key else 10
        last_year = dt.date.today().year
        for a in range(1997, last_year + 1, span):
            body = {"seriesid": ["CUUR0000SEHA"], "startyear": str(a), "endyear": str(min(a + span - 1, last_year))}
            if key:
                body["registrationkey"] = key
            d = json.loads(http_get("https://api.bls.gov/publicAPI/v2/timeseries/data/", json.dumps(body).encode(),
                                    {"Content-Type": "application/json"}))
            if d.get("status") != "REQUEST_SUCCEEDED":
                raise RuntimeError(d.get("message"))
            for o in d["Results"]["series"][0]["data"]:
                if o["period"].startswith("M") and o["period"] != "M13":
                    try:
                        live[pd.Period(f'{o["year"]}-{o["period"][1:]}', "M")] = float(o["value"])
                    except ValueError:
                        pass
        src = ["BLS API v2"]
    except Exception as e:
        log(f"  BLS API unavailable ({str(e)[:90]}); trying FRED")
        live = {}
        fk = os.environ.get("FRED_API_KEY")
        try:
            if not fk:
                raise RuntimeError("no FRED_API_KEY")
            d = json.loads(http_get("https://api.stlouisfed.org/fred/series/observations?series_id=CUUR0000SEHA"
                                    f"&api_key={fk}&file_type=json&observation_start=1997-01-01"))
            for o in d["observations"]:
                if o["value"] not in (".", ""):
                    live[pd.Period(o["date"][:7], "M")] = float(o["value"])
            src = ["FRED API"]
        except Exception as e2:
            log(f"  FRED unavailable ({str(e2)[:60]}); trying DBnomics for history")
            try:
                d = json.loads(http_get("https://api.db.nomics.world/v22/series/BLS/cu/CUUR0000SEHA?observations=1&format=json"))
                doc = d["series"]["docs"][0]
                for p, v in zip(doc["period"], doc["value"]):
                    if v not in ("NA", None):
                        live[pd.Period(p, "M")] = float(v)
                src = ["DBnomics (history) + embedded snapshot (recent)"]
            except Exception as e3:
                log(f"  DBnomics unavailable ({str(e3)[:60]}); using embedded snapshot")
    if live:
        ls = pd.Series(live).sort_index()
        ls = ls[ls.index >= pd.Period(CPI_SNAP_START, "M")]
        idx = s.index.union(ls.index)
        s = s.reindex(idx)
        s.loc[ls.index] = ls.values
    s = s.sort_index()
    s.loc[pd.Period("2025-10", "M")] = np.nan     # uncollected (published value is a carry-forward)
    s = s[s.index >= pd.Period(CPI_SNAP_START, "M")]
    last = s.dropna().index[-1]
    s = s[s.index <= last]
    return s, src[0]


def fetch_zori():
    try:
        raw = http_get("https://files.zillowstatic.com/research/public_csvs/zori/National_zori_uc_sfrcondomfr_sm_month.csv").decode()
        r = list(csv.reader(io.StringIO(raw)))
        hdr, row = r[0], [x for x in r[1:] if x[2] == "United States"][0]
        z = pd.Series([float(v) for v in row[4:] if v != ""], index=pd.PeriodIndex([h[:7] for h, v in zip(hdr[4:], row[4:]) if v != ""], freq="M"))
        return z, "Zillow Research CSV (live)"
    except Exception as e:
        log(f"  Zillow unavailable ({str(e)[:60]}); using embedded ZORI snapshot")
        return snap_series(ZORI_SNAP_START, ZORI_SNAP), "embedded snapshot (2026-10-19)"


def parse_quarter(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year}Q{(v.month - 1) // 3 + 1}"
    s = str(v).strip()
    m = re.match(r"^(\d{4})\s*[-_ ]?\s*[Qq]([1-4])$", s) or re.match(r"^[Qq]([1-4])\s*[-_ ]?\s*(\d{4})$", s)
    if not m:
        return None
    a, b = m.groups()
    return f"{a}Q{b}" if len(a) == 4 else f"{b}Q{a}"


def load_bls_research():
    """Best-effort parse of the BLS R-CPI-NTR/ATR workbook. Returns dict with 'atr' (index level
    Series) and optionally 'ntr', or None. Strict sanity checks; falls back silently to anchors."""
    paths = [os.environ.get("ATR_XLSX", ""), "r-cpi-ntr-and-r-cpi-atr.xlsx",
             os.path.join(SCRIPT_DIR, "data", "r-cpi-ntr-and-r-cpi-atr.xlsx")] + \
        sorted(glob.glob(os.path.join(SCRIPT_DIR, "data", "*.xlsx"))) + glob.glob("/mnt/user-data/uploads/*.xlsx")
    blob = None
    for p in paths:
        if p and os.path.isfile(p):
            blob = open(p, "rb").read()
            log(f"  BLS research workbook: local file {p}")
            break
    if blob is None:
        try:
            blob = http_get("https://www.bls.gov/cpi/research-series/r-cpi-ntr-and-r-cpi-atr.xlsx",
                            headers={"Accept": "*/*", "Referer": "https://www.bls.gov/cpi/research-series/r-cpi-ntr.htm"})
            if not blob.startswith(b"PK"):
                raise RuntimeError("not an xlsx (blocked)")
            log("  BLS research workbook: downloaded")
        except Exception as e:
            log(f"  BLS research workbook not available ({str(e)[:50]}); calibrating to published anchors")
            return None
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
        found = {}
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if not rows:
                continue
            title = (ws.title or "").lower()
            # locate the quarter column and header row
            for ci in range(min(6, max(len(r) for r in rows))):
                qrows = [(i, parse_quarter(r[ci])) for i, r in enumerate(rows) if ci < len(r) and parse_quarter(r[ci])]
                if len(qrows) < 40:
                    continue
                first = qrows[0][0]
                hdr = [" ".join(str(rows[j][k]) for j in range(max(0, first - 4), first)
                                if k < len(rows[j]) and rows[j][k] is not None).lower()
                       for k in range(max(len(r) for r in rows))]
                for k, h in enumerate(hdr):
                    if k == ci:
                        continue
                    vals = {}
                    for i, q in qrows:
                        v = rows[i][k] if k < len(rows[i]) else None
                        if isinstance(v, (int, float)):
                            vals[q] = float(v)
                    if len(vals) < 40:
                        continue
                    ser = pd.Series(vals).sort_index()
                    is_level = ser.min() > 50 and ser.max() < 1000 and "change" not in h and "%" not in h \
                        and "lower" not in h and "upper" not in h and "bound" not in h and "std" not in h
                    if not is_level:
                        continue
                    tag = "atr" if ("atr" in h or "all tenant" in h or "atr" in title or "all tenant" in title) else \
                          ("ntr" if ("ntr" in h or "new tenant" in h or "ntr" in title or "new tenant" in title) else None)
                    if tag and tag not in found:
                        found[tag] = ser
                break
        if "atr" not in found:
            raise RuntimeError("ATR level column not identified")
        atr = found["atr"]
        yoy = 100 * np.log(atr).diff(4)
        if not (atr.index[0] <= "2000Q4" and len(atr) > 90 and yoy.dropna().between(-5, 15).all()):
            raise RuntimeError("ATR sanity check failed")
        log(f"  parsed BLS ATR: {atr.index[0]}..{atr.index[-1]} ({len(atr)} quarters)")
        return found
    except Exception as e:
        log(f"  BLS workbook parse failed ({str(e)[:70]}); calibrating to published anchors")
        return None


# ----------------------------------------------------------------------------------------
# Engine
# ----------------------------------------------------------------------------------------
def kernel(w, span=6):
    p = np.full(6, w / 6.0)
    p[0] += 1.0 - w
    return np.convolve(p, np.ones(span))


def build_problem(cpi, start, end, w, shutdown_aware=True):
    months = pd.period_range(start, end, freq="M")
    pos = {m: i for i, m in enumerate(months)}
    k6, k12 = kernel(w, 6), kernel(w, 12)
    lc = np.log(cpi)
    rows, ys = [], []
    for t in cpi.index:
        if t not in pos or pd.isna(lc.get(t)):
            continue
        prev, k = t - 1, k6
        if shutdown_aware and t == pd.Period("2025-11", "M"):
            prev = t - 2          # Oct-2025 carried forward: Nov change measured from Sep
        if shutdown_aware and t == pd.Period("2026-04", "M"):
            k = k12               # Apr/Oct panel: 1/6 root of a 12-month change
        if prev not in lc.index or pd.isna(lc.get(prev)):
            continue
        i = pos[t]
        if i - (len(k) - 1) < 0:
            continue
        row = np.zeros(len(months))
        row[i - len(k) + 1:i + 1] = k[::-1]
        rows.append(row)
        ys.append(6.0 * (lc[t] - lc[prev]))
    return np.array(rows), np.array(ys), months


def seasonal_design(months):
    S = np.zeros((len(months), 12))
    for i, m in enumerate(months):
        S[i, m.month - 1] = 1.0
    return S - 1.0 / 12.0


def solve(A, y, months, lam, mean_prior=None, mean_prec=None):
    n = len(months)
    S = seasonal_design(months)
    X = np.hstack([A, A @ S])
    D = np.diff(np.eye(n), axis=0)
    P = X.T @ X
    P[:n, :n] += lam * D.T @ D + 1e-8 * np.eye(n)
    P[n:, n:] += 1e-3 * np.eye(12) + 1e-2 * np.ones((12, 12))
    b = X.T @ y
    if mean_prior is not None:
        P[:n, :n] += mean_prec
        b[:n] += mean_prec @ mean_prior
    cov_u = np.linalg.inv(P)
    th = cov_u @ b
    res = y - X @ th
    dof = max(len(y) - np.trace(X @ cov_u @ X.T), 1.0)
    r = float(res @ res / dof)
    return dict(g=th[:n], s=th[n:], cov=cov_u[:n, :n] * r, r=r, S=S, months=months)


def log_marglik(A, y, months, lam):
    n = len(months)
    S = seasonal_design(months)
    X = np.hstack([A, A @ S])
    D = np.diff(np.eye(n), axis=0)
    Q = np.zeros((n + 12, n + 12))
    Q[:n, :n] = lam * D.T @ D + 1e-6 * np.eye(n)
    Q[n:, n:] = 1e-2 * np.ones((12, 12)) + 1e-3 * np.eye(12)
    P = X.T @ X + Q
    th = np.linalg.solve(P, X.T @ y)
    rss = float(np.sum((y - X @ th) ** 2) + th @ Q @ th)
    _, ldP = np.linalg.slogdet(P)
    _, ldQ = np.linalg.slogdet(Q)
    N = len(y)
    return -0.5 * (N * np.log(rss / N) + ldP - ldQ + N)


def quarterly(months, g, cov, alpha_m):
    """Quarterly index = mean of monthly levels; YoY linear in g with analytic sd."""
    n = len(months)
    G = g + alpha_m
    keys = sorted(set((m.qyear, m.quarter) for m in months))
    qp = {k: i for i, k in enumerate(keys)}
    C = np.tril(np.ones((n, n)))
    M = np.zeros((len(keys), n))
    cnt = np.zeros(len(keys))
    for i, m in enumerate(months):
        M[qp[(m.qyear, m.quarter)]] += C[i]
        cnt[qp[(m.qyear, m.quarter)]] += 1
    M /= cnt[:, None]
    qlog = M @ G
    yoy = np.full(len(keys), np.nan)
    sd = np.full(len(keys), np.nan)
    for j in range(4, len(keys)):
        if cnt[j] == 3 and cnt[j - 4] == 3:
            op = M[j] - M[j - 4]
            yoy[j] = 100 * (qlog[j] - qlog[j - 4])
            sd[j] = 100 * math.sqrt(max(op @ cov @ op, 0.0))
    qoq = np.full(len(keys), np.nan)
    qoq[1:] = 400 * np.diff(qlog)
    return pd.DataFrame({"qlog": qlog, "yoy": yoy, "sd": sd, "qoq": qoq, "n": cnt},
                        index=[f"{a}Q{b}" for a, b in keys])


def lease_roll_fit(g_ann, zy, train_end, horizons=range(0, 13), base_lag=4):
    out = {}
    for h in horizons:
        L = max(base_lag, h)
        df = pd.concat([g_ann, zy.shift(L)], axis=1, keys=["g", "z"]).dropna()
        df = df[df.index <= train_end]
        X = np.column_stack([np.ones(len(df)), df["z"].values])
        coef, *_ = np.linalg.lstsq(X, df["g"].values, rcond=None)
        res = df["g"].values - X @ coef
        rho = float(np.corrcoef(res[1:], res[:-1])[0, 1])
        out[h] = dict(a=float(coef[0]), b=float(coef[1]), sd=float(np.std(res, ddof=2)), L=L, n=len(df), rho=rho)
    return out


def calibrate_alpha(months, g, cov, level_targets):
    """alpha_m (monthly) matching the published log level change between two quarters."""
    q0 = quarterly(months, g, cov, 0.0)
    (qa, la), (qb, lb) = level_targets
    nm = 3 * ((int(qb[:4]) - int(qa[:4])) * 4 + int(qb[-1]) - int(qa[-1]))
    return (math.log(lb / la) - (q0.loc[qb, "qlog"] - q0.loc[qa, "qlog"])) / nm


def run_system(cpi, z, w, lam, asof, end, level_targets=None, alpha_m=None, shutdown_aware=True,
               prior_from_lag=PRIOR_FROM_LAG, train_start="2016-01"):
    cpi_t = cpi[cpi.index <= asof]
    z_t = z[z.index <= asof]
    zy = 100 * np.log(z_t).diff(12)
    A, y, months = build_problem(cpi_t, "1996-01", end, w, shutdown_aware)
    s1 = solve(A, y, months, lam)
    am = alpha_m if alpha_m is not None else calibrate_alpha(months, s1["g"], s1["cov"], level_targets)
    g_ann = pd.Series(1200 * (s1["g"] + am), index=months)
    last_cpi = cpi_t.dropna().index[-1]
    train_end = last_cpi - prior_from_lag
    lr = lease_roll_fit(g_ann[g_ann.index >= pd.Period(train_start, "M")], zy, train_end - 4)
    last_z = zy.dropna().index[-1]
    mp = np.zeros(len(months)); sdv = np.zeros(len(months)); idx = []
    for i, m in enumerate(months):
        if m <= train_end:
            continue
        h = max(0, (m - last_z).n)
        if h not in lr:
            continue
        src = m - lr[h]["L"]
        if src not in zy.index or pd.isna(zy.get(src)):
            continue
        mp[i] = (lr[h]["a"] + lr[h]["b"] * zy[src]) / 1200.0 - am
        sdv[i] = lr[h]["sd"] / 1200.0
        idx.append(i)
    idx = np.array(idx, dtype=int)
    rho = min(max(lr[0]["rho"], 0.0), 0.99)
    Pm = np.zeros((len(months), len(months)))
    if len(idx):
        lag = np.abs(idx[:, None] - idx[None, :])
        Sig = np.outer(sdv[idx], sdv[idx]) * rho ** lag
        Pm[np.ix_(idx, idx)] = np.linalg.inv(Sig) * s1["r"]
    s2 = solve(A, y, months, lam, mean_prior=mp, mean_prec=Pm)
    lr_path = pd.Series(np.where(sdv > 0, 1200 * (mp + am), np.nan), index=months)
    return dict(s1=s1, s2=s2, alpha_m=am, months=months, lr=lr, rho=rho, zy=zy, last_cpi=last_cpi,
                train_end=train_end, lr_path=lr_path, A=A, y=y)


def implied_cpi(R, w, cpi):
    """Forward-convolve the estimated reset rate through the CPI estimator."""
    months, s2 = R["months"], R["s2"]
    gs = s2["g"] + s2["S"] @ s2["s"]
    k6 = kernel(w, 6)
    last = R["last_cpi"]
    pos = {m: i for i, m in enumerate(months)}
    fut = [m for m in months if m > last]
    ops = []
    for m in fut:
        row = np.zeros(len(months))
        i = pos[m]
        row[i - len(k6) + 1:i + 1] = k6[::-1] / 6.0
        ops.append(row)
    ops = np.array(ops)
    dl = ops @ gs
    cum_ops = np.cumsum(ops, axis=0)
    lnI = math.log(cpi[last]) + np.cumsum(dl)
    lvl = pd.Series(np.exp(lnI), index=fut)
    sd_cum = np.sqrt(np.maximum(np.einsum("ij,jk,ik->i", cum_ops, s2["cov"], cum_ops), 0))
    full = pd.concat([cpi.dropna(), lvl])
    g0, g1 = pd.Period(GAP_MONTHS[0], "M"), pd.Period(GAP_MONTHS[1], "M")
    out = []
    for j, m in enumerate(fut):
        base = m - 12
        b = full.get(base)
        in_gap = g0 <= base <= g1
        if base == g0:                     # Oct-2025 index not published
            b_u = float(cpi[g0 - 1]) * math.exp(CPI_GAP)
            b = None
        else:
            b_u = None if b is None or pd.isna(b) else float(b) * (math.exp(CPI_GAP) if in_gap else 1.0)
        if b_u is None:
            continue
        yoy = None if b is None or pd.isna(b) else 100 * math.log(lvl[m] / b)
        yoy_u = 100 * math.log(lvl[m] / b_u)
        sd = 100 * sd_cum[j]
        if base > last:   # base itself forecast
            sd = 100 * math.sqrt(max(sd_cum[j] ** 2 - sd_cum[fut.index(base)] ** 2, 0))
        out.append(dict(m=str(m), h=j + 1, yoy=yoy, yoy_u=yoy_u, sd=sd, level=float(lvl[m])))
    return out


# ----------------------------------------------------------------------------------------
# Pipeline
# ----------------------------------------------------------------------------------------
def qkey(p):
    return f"{p.qyear}Q{p.quarter}"


def fnum(x, d=3):
    if x is None:
        return None
    try:
        if np.isnan(x):
            return None
    except TypeError:
        pass
    return round(float(x), d)


class GateError(Exception):
    pass


def sanity_gate(cpi_src, z_src, last_cpi, table, lq, D):
    """Blocks publication of stale or implausible output (used with --strict)."""
    problems = []
    if cpi_src not in LIVE_CPI_SOURCES:
        problems.append(f"CPI rent came from '{cpi_src}', not a live API (set BLS_API_KEY or FRED_API_KEY)")
    if "live" not in z_src:
        problems.append(f"ZORI came from '{z_src}', not the live Zillow file")
    snap_last = pd.Period(CPI_SNAP_START, "M") + (len(CPI_SNAP) - 1)
    if last_cpi < snap_last:
        problems.append(f"latest CPI month {last_cpi} is older than the embedded snapshot ({snap_last})")
    age_days = (dt.date.today() - last_cpi.end_time.date()).days
    if age_days > 75:
        problems.append(f"latest CPI month {last_cpi} is {age_days} days old")
    rows = [r for r in table if r["status"] in ("nowcast", "forecast")]
    if len(rows) < 3 or any(r["yoy"] is None or not (-5 < r["yoy"] < 15) for r in rows):
        problems.append("nowcast/forecast values missing or out of range")
    est = [r["yoy"] for r in table if r["status"] == "estimated" and r["yoy"] is not None]
    if len(est) < 90:
        problems.append(f"only {len(est)} estimated quarters")
    if not D["backtest"] or D["params"]["w"] is None:
        problems.append("backtest or calibration missing")
    if problems:
        raise GateError("; ".join(problems))


def write_archive(D, table, cpi_src):
    sd = os.path.join(SCRIPT_DIR, "state")
    os.makedirs(os.path.join(sd, "vintages"), exist_ok=True)
    tag = f"cpi{D['asof_cpi']}_z{D['asof_z']}"
    pd.DataFrame(table).to_csv(os.path.join(sd, "vintages", f"atri_{tag}.csv"), index=False)
    h = D["headline"]
    rec = dict(built_utc=D["generated"], asof_cpi=D["asof_cpi"], asof_zori=D["asof_z"], cpi_source=cpi_src,
               calibration=D["sources"]["calib"], w=D["params"]["w"], wedge_pp=round(D["params"]["alpha"], 3),
               nowcast_q=h["cur_q"], nowcast_yoy=h["cur"]["yoy"], nowcast_lo=h["cur"]["lo"], nowcast_hi=h["cur"]["hi"],
               forecast_q=h["f_q"], forecast_yoy=h["f"]["yoy"],
               implied_cpi_month=(h["icpi_last"] or {}).get("m"), implied_cpi_yoy=(h["icpi_last"] or {}).get("yoy"))
    logp = os.path.join(sd, "headline_log.csv")
    new = not os.path.exists(logp)
    with open(logp, "a", newline="") as f:
        wr = csv.DictWriter(f, fieldnames=list(rec))
        if new:
            wr.writeheader()
        wr.writerow(rec)
    log(f"  archived vintage {tag}")


def main(argv=None):
    ap = argparse.ArgumentParser(description="Rebuild the Acheron All-Tenant Rent Index dashboard")
    ap.add_argument("--out", default=None, help="output directory")
    ap.add_argument("--html-name", default=HTML_NAME)
    ap.add_argument("--strict", action="store_true", help="require live data and pass the sanity gate")
    ap.add_argument("--archive", action="store_true", help="append headline log and save vintage to state/")
    ap.add_argument("--standalone", action="store_true", help="write a complete HTML document (for GitHub Pages)")
    args = ap.parse_args(argv)
    out_dir = args.out or OUT_DIR
    t0 = dt.datetime.now(dt.timezone.utc)
    log("A-ATRI rebuild " + t0.strftime("%Y-%m-%d %H:%M UTC"))
    cpi, cpi_src = fetch_cpi()
    z, z_src = fetch_zori()
    bls = load_bls_research()
    ntr = bls["ntr"] if bls and "ntr" in bls else pd.Series(NTR_SNAP)
    ntr_src = "BLS workbook" if bls and "ntr" in bls else "BLS via YCharts (2025Q3 vintage)"
    last_cpi = cpi.dropna().index[-1]
    end = last_cpi + H_FCST_M
    log(f"  CPI rent: {cpi_src}, {cpi.index[0]}..{last_cpi}; ZORI: {z_src}, ..{z.index[-1]}")

    # smoothness by marginal likelihood (at a central w)
    A, y, months = build_problem(cpi, "1996-01", str(last_cpi), 0.15)
    lams = np.exp(np.arange(np.log(2.0), np.log(3000.0), 0.2))
    lam = float(lams[int(np.argmax([log_marglik(A, y, months, l) for l in lams]))])
    log(f"  smoothness lambda (max marginal likelihood) = {lam:.1f}")

    full_atr = bls["atr"] if bls else None
    if full_atr is not None:
        mature = full_atr.index[:-4]
        level_targets = ((mature[0], full_atr[mature[0]]), (mature[-1], full_atr[mature[-1]]))
        tgt_yoy = (100 * np.log(full_atr).diff(4))[mature].dropna()
        calib_mode = f"full BLS series ({len(tgt_yoy)} mature quarters)"
    else:
        level_targets = (("1999Q4", 102.275), ("2024Q4", 201.766))
        tgt_yoy = pd.Series({a["q"]: a["value"] for a in ANCHORS if a["kind"] == "yoy" and a["role"] == "w selection"})
        calib_mode = "published anchors (BLS workbook unavailable)"

    # w grid: CPI-only inversion, alpha from level targets, score on target YoY
    wscan = []
    for w in W_GRID:
        A, y, months = build_problem(cpi, "1996-01", str(last_cpi), w)
        s = solve(A, y, months, lam)
        am = calibrate_alpha(months, s["g"], s["cov"], level_targets)
        q = quarterly(months, s["g"], s["cov"], am)
        e = (q["yoy"].reindex(tgt_yoy.index) - tgt_yoy).dropna()
        yy = q["yoy"]["2001Q1":"2024Q4"]
        wscan.append(dict(w=float(w), rmse=float(np.sqrt(np.mean(e ** 2))), alpha=am * 1200,
                          peak=yy.idxmax(), trough=yy.idxmin()))
    best = min(wscan, key=lambda d: d["rmse"])
    W = best["w"]
    log(f"  calibration on {calib_mode}: w = {W:.2f}, RMSE {best['rmse']:.3f}pp, wedge {best['alpha']:+.3f}pp/yr")

    # main run
    R = run_system(cpi, z, W, lam, last_cpi, end, level_targets=level_targets)
    am = R["alpha_m"]
    q1 = quarterly(R["months"], R["s1"]["g"], R["s1"]["cov"], am)
    q2 = quarterly(R["months"], R["s2"]["g"], R["s2"]["cov"], am)
    norm_q, norm_v = level_targets[1]
    lvl = np.exp(q2["qlog"] - q2.loc[norm_q, "qlog"]) * norm_v
    lr_q = R["lr_path"].groupby([R["lr_path"].index.qyear, R["lr_path"].index.quarter]).mean()
    lr_q.index = [f"{a}Q{b}" for a, b in lr_q.index]

    # CPI quarterly YoY (quarter averages; 2025Q4 uses Nov-Dec)
    cq = cpi.groupby([cpi.index.qyear, cpi.index.quarter]).mean()
    cq.index = [f"{a}Q{b}" for a, b in cq.index]
    cpi_q_yoy = 100 * np.log(cq).diff(4)
    cm_yoy = 100 * np.log(cpi).diff(12)

    # real-time backtest (fixed w, alpha; ZORI current vintage)
    bt, cbt = [], []
    for asof in pd.period_range("2018-12", str(last_cpi - 9), freq="Q").asfreq("M", "end"):
        Rb = run_system(cpi, z, W, lam, asof, asof + H_FCST_M, alpha_m=am)
        qb2 = quarterly(Rb["months"], Rb["s2"]["g"], Rb["s2"]["cov"], am)
        qb1 = quarterly(Rb["months"], Rb["s1"]["g"], Rb["s1"]["cov"], am)
        for d in implied_cpi(Rb, W, cpi[cpi.index <= asof]):
            mm = pd.Period(d["m"], "M")
            act = cm_yoy.get(mm)
            gap_hit = pd.Period("2025-10", "M") <= mm <= pd.Period("2026-04", "M") or \
                      pd.Period("2025-10", "M") <= mm - 12 <= pd.Period("2026-04", "M")
            if act is not None and not pd.isna(act) and not gap_hit and d["h"] in (3, 6, 9):
                cbt.append(dict(h=d["h"], e=d["yoy_u"] - act, sd=d["sd"]))
        keys = list(qb2.index)
        i0 = keys.index(qkey(asof))
        for h in range(4):
            if i0 + h >= len(keys):
                break
            k = keys[i0 + h]
            tr = q1["yoy"].get(k)
            if tr is None or pd.isna(tr) or k > qkey(last_cpi - 6):
                continue
            bt.append(dict(asof=str(asof), h=h, q=k, truth=tr, sys=qb2.loc[k, "yoy"], sd=qb2.loc[k, "sd"],
                           cpionly=qb1.loc[k, "yoy"], naive=cm_yoy[asof] + best["alpha"]))
    bt = pd.DataFrame(bt)
    bt_stats = []
    for h, g in bt.groupby("h"):
        rm = lambda c: float(np.sqrt(np.mean((g[c] - g["truth"]) ** 2)))
        k_h = max(1.0, rm("sys") / float(g["sd"].mean()))
        cov90 = float(np.mean(np.abs(g["sys"] - g["truth"]) <= 1.645 * k_h * g["sd"]))
        bt_stats.append(dict(h=int(h), n=len(g), sys=rm("sys"), cpionly=rm("cpionly"), naive=rm("naive"),
                             k=k_h, raw_cov=float(np.mean(np.abs(g["sys"] - g["truth"]) <= 1.645 * g["sd"])), cov=cov90))
    kmap = {d["h"]: d["k"] for d in bt_stats}
    rmap = {d["h"]: d["sys"] for d in bt_stats}
    cbt = pd.DataFrame(cbt)
    cpi_bt = [dict(h=int(h), n=len(g), rmse=float(np.sqrt(np.mean(g["e"] ** 2))), sd=float(g["sd"].mean()))
              for h, g in cbt.groupby("h")] if len(cbt) else []
    log("  implied-CPI backtest RMSE (pp) at +3/+6/+9m: " + ", ".join(f"{d['rmse']:.2f}" for d in cpi_bt))
    log("  backtest RMSE (pp) h0..h3 system: " + ", ".join(f"{d['sys']:.2f}" for d in bt_stats))

    # status & band scaling
    lq = qkey(last_cpi) if last_cpi.month % 3 == 0 else qkey(last_cpi - last_cpi.month % 3)
    qkeys = list(q2.index)
    il = qkeys.index(lq)
    status, scale = {}, {}
    for i, k in enumerate(qkeys):
        if i <= il - 2:
            status[k], scale[k] = "estimated", 0.0
        elif i <= il:
            status[k], scale[k] = "nowcast", rmap.get(0, 0.0)
        else:
            hh = min(i - il, 3)
            status[k], scale[k] = "forecast", rmap.get(hh, rmap[max(rmap)])

    # sensitivity to w
    sens = {}
    for w in [0.0, W, 0.5, 1.0]:
        Rw = run_system(cpi, z, w, lam, last_cpi, end, level_targets=level_targets)
        qw = quarterly(Rw["months"], Rw["s2"]["g"], Rw["s2"]["cov"], Rw["alpha_m"])
        sens[f"{w:.2f}"] = {k: fnum(v) for k, v in qw["yoy"].items() if k >= "2006Q1"}

    # shutdown handling: naive treatment (carry-forward Oct, ordinary Apr kernel)
    cpi_n = cpi.copy()
    cpi_n.loc[pd.Period("2025-10", "M")] = cpi[pd.Period("2025-09", "M")]
    Rn = run_system(cpi_n, z, W, lam, last_cpi, end, alpha_m=am, shutdown_aware=False)
    qn = quarterly(Rn["months"], Rn["s2"]["g"], Rn["s2"]["cov"], am)

    # implied CPI rent path
    icpi = implied_cpi(R, W, cpi)

    # cross-correlations (positive lead = X leads A-ATRI)
    zq = z.groupby([z.index.qyear, z.index.quarter]).mean()
    zq.index = [f"{a}Q{b}" for a, b in zq.index]
    zq_yoy = 100 * np.log(zq).diff(4)
    ntr_yoy = 100 * np.log(ntr).diff(4)
    synth_hist = q1["yoy"][:lq]
    def xcorr(x, lo=-2, hi=6):
        both = pd.concat([synth_hist, x], axis=1, keys=["s", "x"]).dropna()
        return {L: fnum(both["s"].corr(both["x"].shift(L))) for L in range(lo, hi + 1)}, len(both)
    xc = {"ZORI": xcorr(zq_yoy), "BLS NTR": xcorr(ntr_yoy), "CPI rent": xcorr(cpi_q_yoy)}

    # validation table
    vrows = []
    for a in ANCHORS:
        if a["kind"] == "yoy":
            s_ = q1["yoy"].get(a["q"])
            k_next = qkeys[qkeys.index(a["q"]) + 1]
            nv = cpi_q_yoy.get(k_next) + best["alpha"]
        else:
            s_ = float(np.exp(q1.loc[a["q"], "qlog"] - q1.loc[norm_q, "qlog"]) * norm_v)
            nv = None
        vrows.append(dict(q=a["q"], kind=a["kind"], pub=a["value"], src=a["src"], role=a["role"],
                          syn=fnum(s_), err=fnum(s_ - a["value"]), naive=fnum(nv),
                          nerr=fnum(nv - a["value"]) if nv is not None else None))
    yy = q1["yoy"]["2000Q4":"2024Q4"]
    checks = dict(median_pub=ANCHOR_MEDIAN, median_syn=fnum(float(yy.median())),
                  peak_pub=ANCHOR_PEAK, peak_syn=yy.idxmax(), trough_pub=ANCHOR_TROUGH, trough_syn=yy.idxmin(),
                  peak_val=fnum(float(yy.max())), trough_val=fnum(float(yy.min())))

    # quarterly table
    table = []
    for k in qkeys:
        if k < "2000Q1":
            continue
        yv, sd = q2.loc[k, "yoy"], q2.loc[k, "sd"]
        half = 1.645 * max(sd if not pd.isna(sd) else 0.0, scale.get(k, 0.0))
        pub_y = None
        if full_atr is not None and k in full_atr.index:
            pub_y = (100 * np.log(full_atr).diff(4)).get(k)
        else:
            pub_y = next((a["value"] for a in ANCHORS if a["q"] == k and a["kind"] == "yoy"), None)
        table.append(dict(q=k, status=status[k], level=fnum(lvl[k]), yoy=fnum(yv), lo=fnum(yv - half),
                          hi=fnum(yv + half), qoq=fnum(q2.loc[k, "qoq"]), cpi=fnum(cpi_q_yoy.get(k)),
                          cpionly=fnum(q1.loc[k, "yoy"]), leaseroll=None, pub=fnum(pub_y),
                          naive_shutdown=fnum(qn["yoy"].get(k))))
    # lease-roll-only quarterly YoY (where prior exists for all months in window)
    lr_m = R["lr_path"]
    hist_g = pd.Series(1200 * (R["s1"]["g"] + am), index=R["months"])
    mix = hist_g.where(lr_m.isna(), lr_m)
    lv = np.cumsum(mix.values / 1200.0)
    mser = pd.Series(np.exp(lv), index=R["months"])
    mq = mser.groupby([mser.index.qyear, mser.index.quarter]).mean()
    mq.index = [f"{a}Q{b}" for a, b in mq.index]
    mq_yoy = 100 * np.log(mq).diff(4)
    first_lr = lr_m.dropna().index[0]
    for row in table:
        if row["q"] >= qkey(first_lr):
            row["leaseroll"] = fnum(mq_yoy.get(row["q"]))

    # monthly CPI rent YoY (actual) for chart
    cpi_m = [dict(m=str(m), yoy=fnum(v)) for m, v in cm_yoy.items() if m >= pd.Period("2019-01", "M")]

    # headline numbers
    def row(k):
        return next(r for r in table if r["q"] == k)
    cur = row(lq)
    prev_y = row(qkeys[il - 4])
    fq = qkeys[min(il + 3, len(qkeys) - 1)]
    f_row = row(fq)
    last_bls_q = str(full_atr.index[-1]) if full_atr is not None else LAST_BLS_PRINT
    bls_status = "active" if last_bls_q > LAST_BLS_PRINT else "paused"
    last_bls = row(last_bls_q) if any(r["q"] == last_bls_q for r in table) else row(LAST_BLS_PRINT)
    zy_m = 100 * np.log(z).diff(12)
    z_min_m = zy_m["2025-06":].idxmin()
    icpi_last = icpi[-1] if icpi else None
    lead_cpi = max(xc["CPI rent"][0], key=lambda L: xc["CPI rent"][0][L] or -9)

    D = dict(
        generated=t0.strftime("%Y-%m-%d %H:%M UTC"),
        asof_cpi=str(last_cpi), asof_z=str(z.index[-1]), last_q=lq, last_bls=last_bls_q, bls_status=bls_status,
        sources=dict(cpi=cpi_src, zori=z_src, ntr=ntr_src, calib=calib_mode),
        params=dict(w=W, lam=lam, alpha=best["alpha"], rho=R["rho"], prior_from_lag=PRIOR_FROM_LAG,
                    lr0=R["lr"][0], lr9=R["lr"].get(9)),
        table=table, cpi_m=cpi_m, icpi=[{k: fnum(v) if k != "m" else v for k, v in d.items()} for d in icpi],
        cpi_bt=cpi_bt, cpi_gap=CPI_GAP * 100,
        zori_q={k: fnum(v) for k, v in zq_yoy.items() if not pd.isna(v)},
        ntr_q={k: fnum(v) for k, v in ntr_yoy.items() if not pd.isna(v)},
        full_bls={k: fnum(v) for k, v in (100 * np.log(full_atr).diff(4)).items() if not pd.isna(v)} if full_atr is not None else None,
        anchors=[dict(q=a["q"], v=a["value"]) for a in ANCHORS if a["kind"] == "yoy"],
        wscan=wscan, sens=sens, validation=vrows, checks=checks,
        backtest=bt_stats, xcorr={k: dict(c=v[0], n=v[1]) for k, v in xc.items()},
        headline=dict(cur_q=lq, cur=cur, prev=prev_y, f_q=fq, f=f_row, last_bls=last_bls,
                      z_now=fnum(zy_m.iloc[-1], 2), z_now_m=str(zy_m.index[-1]), z_min=fnum(zy_m[z_min_m], 2),
                      z_min_m=str(z_min_m), cpi_now=fnum(cm_yoy.iloc[-1], 2), icpi_last=icpi_last,
                      lead_cpi=int(lead_cpi)),
        log=LOG,
    )
    if args.strict:
        try:
            sanity_gate(cpi_src, z_src, last_cpi, table, lq, D)
            log("  sanity gate passed")
        except GateError as e:
            log(f"  SANITY GATE FAILED - nothing written: {e}")
            sys.exit(2)
    os.makedirs(out_dir, exist_ok=True)
    js = json.dumps(D, separators=(",", ":"), default=lambda o: None if (isinstance(o, float) and np.isnan(o)) else (float(o) if isinstance(o, (np.floating,)) else (int(o) if isinstance(o, np.integer) else str(o))))
    js = js.replace("NaN", "null")
    html, n = re.subn(r"/\*DATA\*/.*?/\*END\*/", lambda m: "/*DATA*/" + js + "/*END*/", TEMPLATE, flags=re.S)
    assert n == 1, "data placeholder not found"
    if args.standalone:
        cut = html.index("</style>") + len("</style>")
        html = ('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
                '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
                '<style>[hidden]{display:none!important}img{max-width:100%}:root{padding-top:env(safe-area-inset-top,0px);'
                'padding-bottom:env(safe-area-inset-bottom,0px)}</style>\n'
                + html[:cut] + "\n</head>\n<body>\n" + html[cut:] + "\n</body>\n</html>\n")
    with open(os.path.join(out_dir, args.html_name), "w", encoding="utf-8") as f:
        f.write(html)
    pd.DataFrame(table).to_csv(os.path.join(out_dir, CSV_NAME), index=False)
    latest = dict(generated=D["generated"], asof_cpi=D["asof_cpi"], asof_zori=D["asof_z"],
                  bls_status=bls_status, last_bls_print=last_bls_q, sources=D["sources"],
                  nowcast=dict(q=lq, yoy=cur["yoy"], lo=cur["lo"], hi=cur["hi"]),
                  forecast=dict(q=fq, yoy=f_row["yoy"], lo=f_row["lo"], hi=f_row["hi"]),
                  implied_cpi_rent=icpi_last, wedge_pp_per_year=round(best["alpha"], 3), redating_w=W)
    with open(os.path.join(out_dir, "latest.json"), "w") as f:
        json.dump(latest, f, indent=2, default=lambda o: None)
    if args.html_name == "index.html":
        open(os.path.join(out_dir, ".nojekyll"), "w").close()
    log(f"  wrote {os.path.join(out_dir, args.html_name)}, {CSV_NAME}, latest.json")
    if args.archive:
        write_archive(D, table, cpi_src)
    log(f"  headline: {lq} A-ATRI {cur['yoy']:.2f}% y/y [{cur['lo']:.2f}, {cur['hi']:.2f}]; {fq} forecast {f_row['yoy']:.2f}%")


TEMPLATE = r"""<title>Acheron All-Tenant Rent Index</title>
<meta name="claude:viewport-fit-sides" content="cover">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;500;600&family=IBM+Plex+Sans:wght@400;500;600&family=Space+Grotesk:wght@500;600;700&display=swap">
<style>
/* Layout: research-note column with an instrument strip; charts sit in cards, prose stays on paper */
:root{
  --paper:#F7F1E6; --card:#FDFAF3; --ink:#221C14; --ink2:#554B40; --muted:#857A6B; --rule:#E3D8C6; --grid:#ECE3D4;
  --orange:#D2622A; --orange-lt:#E9A07A; --orange-dk:#9A4317; --orange-xdk:#5E2508;
  --teal:#0E756C; --ochre:#A67A22; --plum:#7A4A66; --slate:#4E6577; --neutral:#A39886;
  --chip-now:#F3E6C8; --chip-fc:#F8E1D3; --chip-paused:#EEDFE8;
  --f-display:"Space Grotesk", "Helvetica Neue", Arial, sans-serif;
  --f-body:"IBM Plex Sans", "Helvetica Neue", Arial, sans-serif;
  --f-mono:"IBM Plex Mono", ui-monospace, "SFMono-Regular", Menlo, monospace;
}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){
  --paper:#17130F; --card:#211C17; --ink:#F0E8DB; --ink2:#C9BDAC; --muted:#958A7B; --rule:#3A3229; --grid:#2E2720;
  --orange:#E07A45; --orange-lt:#F2B08A; --orange-dk:#B85A28; --orange-xdk:#8A3C14;
  --teal:#2BA394; --ochre:#CDA046; --plum:#BC86AC; --slate:#8BA1B6; --neutral:#7A7064;
  --chip-now:#3B3120; --chip-fc:#3E2719; --chip-paused:#352631; color-scheme:dark}}
:root[data-theme="dark"]{
  --paper:#17130F; --card:#211C17; --ink:#F0E8DB; --ink2:#C9BDAC; --muted:#958A7B; --rule:#3A3229; --grid:#2E2720;
  --orange:#E07A45; --orange-lt:#F2B08A; --orange-dk:#B85A28; --orange-xdk:#8A3C14;
  --teal:#2BA394; --ochre:#CDA046; --plum:#BC86AC; --slate:#8BA1B6; --neutral:#7A7064;
  --chip-now:#3B3120; --chip-fc:#3E2719; --chip-paused:#352631; color-scheme:dark}
*{box-sizing:border-box}
body{background:var(--paper);color:var(--ink);font-family:var(--f-body);font-size:15px;line-height:1.55;margin:0}
.wrap{max-width:1180px;margin:0 auto;padding-left:max(16px, env(safe-area-inset-left, 0px));padding-right:max(16px, env(safe-area-inset-right, 0px));padding-block:28px 64px;display:flex;flex-direction:column;gap:40px}
h1,h2,h3{font-family:var(--f-display);text-wrap:balance;margin:0;color:var(--ink)}
h1{font-size:clamp(30px,4.4vw,46px);font-weight:600;letter-spacing:-.015em;line-height:1.08}
h2{font-size:22px;font-weight:600;letter-spacing:-.005em}
h3{font-size:15px;font-weight:600}
p{margin:0;max-width:72ch}
a{color:var(--orange-dk)}
:root:not([data-theme="light"]) a{color:var(--orange)}
.eyebrow,.lab{font-family:var(--f-mono);font-size:11.5px;letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}
.mono,.num{font-family:var(--f-mono);font-variant-numeric:tabular-nums}
header.mast{display:flex;flex-direction:column;gap:14px;border-bottom:1px solid var(--rule);padding-bottom:24px}
.dek{font-size:17px;color:var(--ink2);max-width:68ch}
.meta{display:flex;flex-wrap:wrap;gap:8px 18px;font-family:var(--f-mono);font-size:12px;color:var(--ink2)}
.meta span b{color:var(--ink);font-weight:500}
nav.toc{display:flex;flex-wrap:wrap;gap:6px}
nav.toc a{font-family:var(--f-mono);font-size:12px;text-decoration:none;color:var(--ink2);border:1px solid var(--rule);border-radius:999px;padding:4px 11px;background:var(--card)}
nav.toc a:hover,nav.toc a:focus-visible{border-color:var(--orange);color:var(--ink);outline:none}
.kpis{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px}
.kpi{background:var(--card);border:1px solid var(--rule);border-radius:6px;padding:16px 16px 14px;display:flex;flex-direction:column;gap:6px;min-width:0}
.kpi .big{font-family:var(--f-mono);font-size:34px;font-weight:500;line-height:1;letter-spacing:-.02em;font-variant-numeric:tabular-nums}
.kpi .big small{font-size:16px;color:var(--ink2);margin-left:2px}
.kpi .sub{font-size:13px;color:var(--ink2)}
.kpi.hero{border-top:3px solid var(--orange)}
.chip{display:inline-flex;align-items:center;gap:5px;font-family:var(--f-mono);font-size:11px;letter-spacing:.04em;text-transform:uppercase;border-radius:3px;padding:2px 7px;border:1px solid transparent;white-space:nowrap}
.chip.estimated{border-color:var(--rule);color:var(--ink2)}
.chip.nowcast{background:var(--chip-now);color:var(--ink)}
.chip.forecast{background:var(--chip-fc);color:var(--ink);border:1px dashed var(--orange)}
.chip.paused{background:var(--chip-paused);color:var(--ink)}
.chip::before{content:"";width:6px;height:6px;border-radius:50%;background:currentColor;opacity:.55}
.chip.nowcast::before{background:var(--ochre);opacity:1}.chip.forecast::before{background:var(--orange);opacity:1}.chip.paused::before{background:var(--plum);opacity:1}
.findings{display:grid;grid-template-columns:minmax(0,1.25fr) minmax(0,1fr);gap:28px;align-items:start}
.findings ul{margin:0;padding-left:18px;display:flex;flex-direction:column;gap:8px}
.findings li{max-width:70ch}
section{display:flex;flex-direction:column;gap:14px;scroll-margin-top:16px}
.sec-head{display:flex;flex-direction:column;gap:6px}
.sec-head p{color:var(--ink2)}
.panel{background:var(--card);border:1px solid var(--rule);border-radius:6px;padding:16px;display:flex;flex-direction:column;gap:10px;min-width:0}
.panel-top{display:flex;justify-content:space-between;align-items:baseline;gap:10px;flex-wrap:wrap}
.chart{position:relative;height:340px;width:100%}
.chart.sm{height:260px}
.grid2{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px}
.legend{display:flex;flex-wrap:wrap;gap:6px 16px;font-size:12.5px;color:var(--ink2)}
.legend i{display:inline-block;width:18px;height:0;border-top:2.5px solid;vertical-align:middle;margin-right:6px}
.legend i.dash{border-top-style:dashed}.legend i.dot{border-top-style:dotted}
.legend i.pt{width:9px;height:9px;border:2px solid var(--card);border-radius:50%;box-shadow:0 0 0 1px currentColor}
.legend i.band{height:10px;border:none;border-radius:2px}
.seg{display:inline-flex;border:1px solid var(--rule);border-radius:4px;overflow:hidden}
.seg button{font-family:var(--f-mono);font-size:12px;background:transparent;border:0;padding:5px 10px;color:var(--ink2);cursor:pointer}
.seg button[aria-pressed="true"]{background:var(--ink);color:var(--paper)}
.seg button:focus-visible{outline:2px solid var(--orange);outline-offset:-2px}
.note{font-size:12.5px;color:var(--muted);max-width:none}
.tbl-wrap{overflow-x:auto;border:1px solid var(--rule);border-radius:6px;background:var(--card)}
table{border-collapse:collapse;width:100%;font-size:13px}
th,td{padding:7px 10px;text-align:right;border-bottom:1px solid var(--grid);white-space:nowrap}
th{font-family:var(--f-mono);font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:var(--muted);font-weight:500;background:var(--card);position:sticky;top:0}
td{font-family:var(--f-mono);font-variant-numeric:tabular-nums}
th:first-child,td:first-child{text-align:left}
td.txt,th.txt{text-align:left;font-family:var(--f-body);white-space:normal;min-width:180px}
tr:last-child td{border-bottom:0}
td.hl{color:var(--orange-dk);font-weight:600}
:root:not([data-theme="light"]) td.hl{color:var(--orange)}
.pos{color:var(--ink)} .neg{color:var(--ink)}
.btn{font-family:var(--f-mono);font-size:12px;border:1px solid var(--ink);background:transparent;color:var(--ink);border-radius:4px;padding:6px 12px;cursor:pointer}
.btn:hover,.btn:focus-visible{background:var(--ink);color:var(--paper);outline:none}
.method{display:grid;grid-template-columns:minmax(0,1fr) minmax(0,1fr);gap:24px}
.eq{font-family:var(--f-mono);font-size:13px;background:var(--card);border:1px solid var(--rule);border-left:3px solid var(--slate);padding:10px 12px;border-radius:4px;overflow-x:auto;white-space:pre;line-height:1.6}
.method ol,.method ul{margin:0;padding-left:20px;display:flex;flex-direction:column;gap:8px}
.method li{max-width:70ch}
.src{font-size:13px;color:var(--ink2)}
.src li{max-width:none}
footer{border-top:1px solid var(--rule);padding-top:14px;font-size:12px;color:var(--muted);font-family:var(--f-mono)}
@media (max-width:900px){.kpis{grid-template-columns:repeat(2,minmax(0,1fr))}.findings,.method,.grid2{grid-template-columns:minmax(0,1fr)}}
@media (max-width:480px){.kpis{grid-template-columns:minmax(0,1fr)}.chart{height:270px}.chart.sm{height:230px}.kpi .big{font-size:30px}}
@media (prefers-reduced-motion:reduce){*{transition:none!important;animation:none!important}}
</style>

<div class="wrap">
  <header class="mast">
    <div class="eyebrow">Acheron Insights · US shelter inflation · synthetic research series</div>
    <h1>Acheron All-Tenant Rent Index</h1>
    <p class="dek">A continuation of the BLS research index R-CPI-ATR, which BLS stopped publishing in April 2026 after the October 2025 housing survey went uncollected. Rebuilt from public CPI rent data by inverting the CPI estimator and re-dating each rent change to the month it was set, then extended with new-lease rents.</p>
    <div class="meta" id="meta"></div>
    <nav class="toc" aria-label="Sections">
      <a href="#index">Index</a><a href="#edge">Nowcast &amp; forecast</a><a href="#leads">Lead indicators</a><a href="#validation">Validation</a><a href="#shutdown">Shutdown gap</a><a href="#cpi">CPI rent outlook</a><a href="#method">Method</a><a href="#data">Data</a>
    </nav>
  </header>

  <div class="kpis" id="kpis"></div>

  <div class="findings">
    <div class="sec-head"><span class="lab">Reading</span><ul id="findings"></ul></div>
    <div class="sec-head"><span class="lab">Status of each quarter</span>
      <p class="src"><span class="chip estimated">Estimated</span> CPI has fully revealed the rent changes set in the quarter.<br>
      <span class="chip nowcast">Nowcast</span> The last two CPI quarters: some rent changes set in the quarter are not yet in the CPI sample, so new-lease rents carry part of the weight.<br>
      <span class="chip forecast">Forecast</span> Beyond the CPI data; driven by Zillow new-lease rents through the lease-roll model. Bands are at least as wide as the backtest error.<br>
      <span class="chip paused">BLS paused</span> No official print after <span class="num" id="lastbls"></span>.</p>
    </div>
  </div>

  <section id="index">
    <div class="sec-head"><span class="lab">01 · The index</span><h2>A-ATRI year-over-year, with the CPI rent it leads</h2>
      <p>Quarterly, four-quarter log change. The shaded band is the 90% interval. Dots are the published BLS prints used to calibrate and test the replication.</p></div>
    <div class="panel">
      <div class="panel-top">
        <div class="legend" id="lg-main"></div>
        <div class="seg" role="group" aria-label="Range"><button id="r-full" aria-pressed="false">2000–</button><button id="r-recent" aria-pressed="true">2018–</button></div>
      </div>
      <div class="chart"><canvas id="c-main" aria-label="A-ATRI year-over-year chart" role="img"></canvas></div>
    </div>
  </section>

  <section id="edge">
    <div class="sec-head"><span class="lab">02 · Nowcast &amp; forecast</span><h2>Where the CPI stops and new-lease rents take over</h2>
      <p id="edge-text"></p></div>
    <div class="panel">
      <div class="legend" id="lg-edge"></div>
      <div class="chart"><canvas id="c-edge" aria-label="Ragged edge chart" role="img"></canvas></div>
      <p class="note" id="edge-note"></p>
    </div>
  </section>

  <section id="leads">
    <div class="sec-head"><span class="lab">03 · Lead indicators</span><h2>New-lease rents lead the all-tenant index by three quarters</h2>
      <p>Both series are shifted forward three quarters. Adams, Loewenstein, Montag and Verbrugge (2023) find the BLS new-tenant index leads the all-tenant index by about three quarters; the replication reproduces that lead with BLS's own NTR and with Zillow's ZORI.</p></div>
    <div class="grid2">
      <div class="panel"><div class="panel-top"><h3>Zillow ZORI, shifted +3Q</h3><span class="num note" id="zc"></span></div><div class="legend" id="lg-z"></div><div class="chart sm"><canvas id="c-z" aria-label="ZORI lead chart" role="img"></canvas></div></div>
      <div class="panel"><div class="panel-top"><h3>BLS new-tenant index, shifted +3Q</h3><span class="num note" id="nc"></span></div><div class="legend" id="lg-n"></div><div class="chart sm"><canvas id="c-n" aria-label="NTR lead chart" role="img"></canvas></div></div>
    </div>
    <div class="tbl-wrap"><table id="t-xc"></table></div>
    <p class="note">Correlation of quarterly YoY with A-ATRI, by lead in quarters (positive: the series leads A-ATRI). Highlighted cell is the peak. A-ATRI history only (estimated quarters).</p>
  </section>

  <section id="validation">
    <div class="sec-head"><span class="lab">04 · Validation</span><h2>Calibration against published BLS prints</h2>
      <p id="val-text"></p></div>
    <div class="tbl-wrap"><table id="t-val"></table></div>
    <p class="note" id="val-checks"></p>
    <div class="grid2">
      <div class="panel"><div class="panel-top"><h3>Anchor fit by re-dating weight w</h3><span class="note">RMSE, pp</span></div><div class="chart sm"><canvas id="c-wscan" aria-label="Fit by re-dating weight" role="img"></canvas></div>
        <p class="note">w = 0: CPI's sixth-root smoothing only. w = 1: every rent change moved back to its lease anniversary, as the BLS dating rule implies in theory. Full re-dating predicts a two-quarter lead over CPI; the published prints show one.</p></div>
      <div class="panel"><div class="legend" id="lg-sens"></div><div class="chart sm"><canvas id="c-sens" aria-label="Sensitivity to w" role="img"></canvas></div>
        <p class="note">A-ATRI under alternative re-dating weights, each re-calibrated. The choice moves turning points by about a quarter and barely moves the current level.</p></div>
    </div>
    <div class="grid2">
      <div class="panel"><div class="panel-top"><h3>Real-time backtest</h3><span class="note">RMSE vs final estimate, pp</span></div><div class="legend" id="lg-bt"></div><div class="chart sm"><canvas id="c-bt" aria-label="Backtest RMSE" role="img"></canvas></div></div>
      <div class="panel"><h3>What the backtest does</h3><p class="note" style="font-size:13.5px;color:var(--ink2)" id="bt-text"></p><div class="tbl-wrap"><table id="t-bt"></table></div></div>
    </div>
  </section>

  <section id="shutdown">
    <div class="sec-head"><span class="lab">05 · Shutdown gap</span><h2>Handling October 2025 the way BLS computed it</h2>
      <p>BLS carried October 2025 rents forward (zero change), then measured the April 2026 change as the sixth root of a twelve-month change for that panel. A-ATRI treats October as unobserved and gives April its twelve-month kernel. Treating the published series at face value instead understates the index through the gap and overshoots after the April catch-up.</p></div>
    <div class="panel"><div class="legend" id="lg-sd"></div><div class="chart sm"><canvas id="c-sd" aria-label="Shutdown treatment comparison" role="img"></canvas></div></div>
  </section>

  <section id="cpi">
    <div class="sec-head"><span class="lab">06 · CPI rent outlook</span><h2>Implied path for CPI rent of primary residence</h2>
      <p id="cpi-text"></p></div>
    <div class="panel"><div class="legend" id="lg-cpi"></div><div class="chart"><canvas id="c-cpi" aria-label="Implied CPI rent path" role="img"></canvas></div>
      <p class="note">Monthly, twelve-month change, NSA. October 2025 is missing in the official series, and published changes from November 2025 to March 2026 are understated by the carry-forward. The implied path forward-convolves the A-ATRI reset rate (before the formula wedge) through CPI's six-month panel estimator; the underlying line corrects the base months using BLS's interpolation counterfactual (Monthly Labor Review, 2026).</p></div>
  </section>

  <section id="method">
    <div class="sec-head"><span class="lab">07 · Method</span><h2>How the replication works</h2></div>
    <div class="method">
      <div class="sec-head" style="gap:12px">
        <p>BLS built R-CPI-ATR from CPI housing-survey microdata as a Case–Shiller repeat-rent regression on consecutive six-month observations of each unit, dating every observation to the tenant's move-in month or the latest six-month multiple after it. Without microdata, the same object can be recovered from the published CPI rent index, because CPI rent is a known function of the rent changes ATR measures.</p>
        <div class="eq">CPI estimator     Δln I_t = ⅙ · ln[ Σ wᵢ r*ᵢ(t) / Σ wᵢ r*ᵢ(t−6) ]
Inversion         6·Δln I_t = Σₖ κₖ · g_{t−k} + ε_t
Kernel            κ = p ∗ 𝟙[0..5],  p = w·U{0..5} + (1−w)·δ₀
Smoothness        Δg_t ~ N(0, σ²/λ),  λ by marginal likelihood
Oct-2025          unobserved; Nov-2025 measured from Sep-2025
Apr-2026          row uses κ₁₂ = p ∗ 𝟙[0..11]
Formula wedge     A-ATRI growth = g_t + α
Lease-roll prior  g_t = a_h + b_h · ZORIyoy_{t−max(4,h)} + u_t,  u AR(1)
Aggregation       quarter = mean of monthly levels; YoY = 4-quarter log change</div>
        <ol>
          <li><b>Inversion.</b> Each month's CPI rent change is one sixth of the six-month change of the panel surveyed that month. Multiplying by six recovers the six-month change, and the kernel maps it onto rent changes dated to when they were set. A Bayesian deconvolution with a random-walk prior on growth recovers the monthly reset rate g with its posterior covariance.</li>
          <li><b>Formula wedge.</b> ATR averages log changes (geometric); CPI rent is a ratio of arithmetic means, which runs higher by roughly half the cross-unit variance of rent changes. α is set so the replication matches the published 1999Q4 to 2024Q4 level ratio exactly.</li>
          <li><b>Re-dating weight.</b> w is chosen on the published YoY prints (see Validation). Seasonal effects are estimated alongside and excluded from the headline, which is seasonally smooth.</li>
          <li><b>Ragged edge and forecast.</b> Direct h-step regressions of the reset rate on lagged ZORI growth form a prior with AR(1)-correlated errors over the last four CPI months and the forecast months. Nowcast and forecast bands are the larger of posterior uncertainty and the real-time backtest RMSE at that horizon.</li>
        </ol>
      </div>
      <div class="sec-head" style="gap:12px">
        <h3>Limitations</h3>
        <ul>
          <li>No microdata. Differences from ATR that do not show up in the CPI aggregate (outlier trimming, remodel exclusions, vacancy rules, BLS's April 2025 switch to expenditure weights) are absorbed into the constant wedge.</li>
          <li id="lim-calib">Calibration uses a handful of published prints from the pre-2025 vintage, not the full series, because BLS blocks automated download of its workbook. Rebuild with the workbook alongside and the script re-calibrates on every quarter automatically.</li>
          <li>The wedge is treated as constant. The published peak and trough sit slightly further apart than the replication, consistent with a wedge that widens when rent changes are more dispersed.</li>
          <li>The backtest compares against the replication's own final estimates and uses the current ZORI vintage, so it measures edge revision rather than error against BLS.</li>
          <li>Apartment List national data could not be retrieved as a static file and is not used.</li>
        </ul>
        <h3>Sources</h3>
        <ul class="src" id="sources"></ul>
      </div>
    </div>
  </section>

  <section id="data">
    <div class="sec-head"><span class="lab">08 · Data</span><h2>Quarterly table</h2></div>
    <div class="panel-top"><div class="seg" role="group" aria-label="Rows"><button id="d-recent" aria-pressed="true">2018–</button><button id="d-all" aria-pressed="false">All</button></div><button class="btn" id="copycsv">Copy CSV</button></div>
    <div class="tbl-wrap" style="max-height:520px"><table id="t-data"></table></div>
    <p class="note" id="copymsg">Index level on the BLS R-CPI-ATR scale (matched at 2024Q4). q/q SAAR is the annualized quarterly change of the seasonally smooth index.</p>
  </section>

  <footer id="foot"></footer>
</div>

<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<script>
const D = /*DATA*/null/*END*/;
(function(){
"use strict";
const $=id=>document.getElementById(id);
const tok=n=>getComputedStyle(document.documentElement).getPropertyValue(n).trim();
function rgba(hex,a){let h=hex.replace('#','');if(h.length===3)h=h.split('').map(c=>c+c).join('');const n=parseInt(h,16);return `rgba(${(n>>16)&255},${(n>>8)&255},${n&255},${a})`;}
const f2=v=>v==null||isNaN(v)?'–':(+v).toFixed(2);
const sg=v=>v==null||isNaN(v)?'–':((v>0?'+':'')+(+v).toFixed(2));
const pct=v=>v==null?'–':f2(v)+'%';
const T=D.table, byQ={}; T.forEach(r=>byQ[r.q]=r);
const H=D.headline, P=D.params;
const mlab=m=>{const [y,mm]=m.split('-');return ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'][+mm-1]+' '+y;};

// ---------- header, KPIs, findings ----------
$('meta').innerHTML=[
  ['CPI rent through',mlab(D.asof_cpi)],['ZORI through',mlab(D.asof_z)],['Last BLS print',D.last_bls],
  ['Calibration',D.sources.calib],['Generated',D.generated]].map(([a,b])=>`<span>${a} <b>${b}</b></span>`).join('');
$('lastbls').textContent=D.last_bls;
const cur=H.cur, fr=H.f;
const kpis=[
  {cls:'hero',lab:`${H.cur_q} · nowcast`,big:f2(cur.yoy),sub:`90% band ${f2(cur.lo)} to ${f2(cur.hi)} · year ago ${f2(H.prev.yoy)}`,chip:'nowcast'},
  {lab:`${H.f_q} · forecast`,big:f2(fr.yoy),sub:`90% band ${f2(fr.lo)} to ${f2(fr.hi)}`,chip:'forecast'},
  {lab:`CPI rent · ${mlab(D.asof_cpi)}`,big:f2(H.cpi_now),sub:H.icpi_last?`Implied ${mlab(H.icpi_last.m)}: ${f2(H.icpi_last.yoy)}%`:'',chip:null},
  D.bls_status==='paused'
    ? {lab:`BLS R-CPI-ATR`,big:'Paused',sub:`Last print ${D.last_bls}, released 29 Jan 2026. Publication paused April 2026.`,chip:'paused',word:true}
    : {lab:`BLS R-CPI-ATR · ${D.last_bls}`,big:f2(H.last_bls.yoy),sub:`BLS has resumed publishing. A-ATRI is now calibrated to the full series.`,chip:null}
];
$('kpis').innerHTML=kpis.map(k=>`<div class="kpi ${k.cls||''}"><div class="panel-top"><span class="lab">${k.lab}</span>${k.chip?`<span class="chip ${k.chip}">${k.chip==='paused'?'BLS paused':k.chip}</span>`:''}</div><div class="big">${k.big}${k.word?'':'<small>% y/y</small>'}</div><div class="sub">${k.sub}</div></div>`).join('');
const dirNow = cur.yoy < H.prev.yoy ? 'down from' : 'up from';
const prevTxt = H.prev.q===D.last_bls ? `${f2(H.prev.yoy)}% a year earlier in ${D.last_bls}, the last quarter BLS published` : `${f2(H.prev.yoy)}% a year earlier and ${f2(H.last_bls.yoy)}% in ${D.last_bls}, the last quarter BLS published`;
const icp=D.icpi.filter(d=>d.yoy!=null); const hump=icp.length?icp.reduce((a,b)=>b.yoy-b.yoy_u>a.yoy-a.yoy_u?b:a):null;
const zTurn = H.z_now > H.z_min + 0.15;
$('findings').innerHTML=[
  `All-tenant rent inflation is running at <b class="num">${f2(cur.yoy)}%</b> y/y in ${H.cur_q} (90% band ${f2(cur.lo)}–${f2(cur.hi)}), ${dirNow} ${prevTxt}.`,
  zTurn ? `New-lease rents have turned. Zillow's ZORI rose from ${f2(H.z_min)}% y/y in ${mlab(H.z_min_m)} to ${f2(H.z_now)}% in ${mlab(H.z_now_m)}; through the lease-roll lag, A-ATRI is forecast at <b class="num">${f2(fr.yoy)}%</b> by ${H.f_q}.` : `New-lease rents are flat: ZORI is ${f2(H.z_now)}% y/y in ${mlab(H.z_now_m)}. A-ATRI is forecast at <b class="num">${f2(fr.yoy)}%</b> by ${H.f_q}.`,
  H.icpi_last ? `Because CPI rent lags A-ATRI by about a quarter, the implied CPI rent path is <b class="num">${f2(H.icpi_last.yoy)}%</b> y/y by ${mlab(H.icpi_last.m)}, against ${f2(H.cpi_now)}% now.`+(hump&&hump.yoy-hump.yoy_u>0.1?` Published prints for Nov 2026 to Mar 2027 will run about ${f2(D.cpi_gap)}pp above the underlying rate, a base effect from the October 2025 carry-forward that unwinds in April 2027.`:'') : '',
  `Mechanically, the BLS index is CPI rent pulled forward about one quarter, less a formula wedge of <b class="num">${f2(-P.alpha)}pp</b> a year. Full lease-anniversary re-dating, which implies a two-quarter lead, is rejected by the published prints; the best fit uses w = ${P.w.toFixed(2)}.`
].filter(Boolean).map(s=>`<li>${s}</li>`).join('');

// ---------- chart plumbing ----------
const charts={};
function palette(){return {ink:tok('--ink'),ink2:tok('--ink2'),muted:tok('--muted'),grid:tok('--grid'),card:tok('--card'),paper:tok('--paper'),
  orange:tok('--orange'),orangeLt:tok('--orange-lt'),orangeDk:tok('--orange-dk'),orangeXdk:tok('--orange-xdk'),teal:tok('--teal'),plum:tok('--plum'),slate:tok('--slate'),ochre:tok('--ochre'),neutral:tok('--neutral'),chipFc:tok('--chip-fc')};}
const shade={id:'shade',beforeDatasetsDraw(c,a,o){if(o.from==null||o.from<0)return;const {ctx,chartArea:ar,scales:{x}}=c;const p=x.getPixelForValue(o.from),q=o.from>0?x.getPixelForValue(o.from-1):p-10;const x0=(p+q)/2;ctx.save();ctx.fillStyle=o.color;ctx.fillRect(x0,ar.top,ar.right-x0,ar.bottom-ar.top);if(o.label){ctx.fillStyle=o.text;ctx.font='11px "IBM Plex Mono", monospace';ctx.textAlign='left';ctx.fillText(o.label,x0+6,ar.top+13);}ctx.restore();}};
const cross={id:'cross',afterDatasetsDraw(c){const t=c.tooltip;if(!t||!t.getActiveElements().length)return;const x=t.getActiveElements()[0].element.x;const {ctx,chartArea:a}=c;ctx.save();ctx.strokeStyle=c.options.plugins.cross.color;ctx.lineWidth=1;ctx.setLineDash([3,3]);ctx.beginPath();ctx.moveTo(x,a.top);ctx.lineTo(x,a.bottom);ctx.stroke();ctx.restore();}};
const zero={id:'zero',beforeDatasetsDraw(c,a,o){const y=c.scales.y;if(!y||y.min>0||y.max<0)return;const {ctx,chartArea:ar}=c;const py=y.getPixelForValue(0);ctx.save();ctx.strokeStyle=o.color;ctx.lineWidth=1;ctx.beginPath();ctx.moveTo(ar.left,py);ctx.lineTo(ar.right,py);ctx.stroke();ctx.restore();}};
Chart.register(shade,cross,zero);
function baseOpts(C,labels,{yTitle='% y/y',step=null,shadeFrom=null,shadeLabel=null,tipFmt=null}={}){
  const n=labels.length; const yrStep=step||(n>90?4:n>44?2:1);
  return {responsive:true,maintainAspectRatio:false,animation:false,interaction:{mode:'index',intersect:false},
    layout:{padding:{top:4,right:6}},
    plugins:{legend:{display:false},cross:{color:C.muted},zero:{color:C.muted},
      shade:{from:shadeFrom,color:rgba(C.orange,0.07),label:shadeLabel,text:C.muted},
      tooltip:{backgroundColor:C.card,titleColor:C.ink,bodyColor:C.ink2,borderColor:C.grid,borderWidth:1,padding:9,
        titleFont:{family:'IBM Plex Mono',size:12},bodyFont:{family:'IBM Plex Mono',size:12},boxPadding:4,usePointStyle:true,
        filter:i=>!i.dataset._band&&i.raw!=null,
        callbacks:{label:tipFmt||(i=>` ${i.dataset.label}: ${f2(i.raw)}%`)}}},
    scales:{x:{grid:{display:false},border:{color:C.grid},ticks:{color:C.muted,font:{family:'IBM Plex Mono',size:11},autoSkip:false,maxRotation:0,
        callback(v){const l=labels[v];if(!l)return '';if(l.includes('Q')){return (l.endsWith('Q1')&&(+l.slice(0,4))%yrStep===0)?l.slice(0,4):'';}
          return (l.endsWith('-01')&&(+l.slice(0,4))%yrStep===0)?l.slice(0,4):'';}}},
      y:{grid:{color:C.grid},border:{display:false},ticks:{color:C.muted,font:{family:'IBM Plex Mono',size:11},callback:v=>v+'%'},title:{display:false,text:yTitle}}}};
}
const line=(label,data,color,o={})=>Object.assign({label,data,borderColor:color,backgroundColor:color,borderWidth:2,pointRadius:0,pointHoverRadius:4,tension:0.25,spanGaps:false},o);
const band=(up,lo,color,a=0.16)=>[{label:'band hi',data:up,borderWidth:0,pointRadius:0,fill:false,_band:true,tension:0.25},{label:'band lo',data:lo,borderWidth:0,pointRadius:0,fill:'-1',backgroundColor:rgba(color,a),_band:true,tension:0.25}];
const pts=(label,data,C)=>({label,data,showLine:false,pointRadius:5,pointHoverRadius:6,pointBackgroundColor:C.ink,pointBorderColor:C.card,pointBorderWidth:2,borderColor:C.ink,backgroundColor:C.ink});
function legend(id,items){$(id).innerHTML=items.map(([t,c,k])=>{const st=k==='band'?`background:${c}`:(k==='pt'?`color:${c};background:${c}`:`border-top-color:${c}`);return `<span><i class="${k||''}" style="${st}"></i>${t}</span>`;}).join('');}
function make(id,cfg){if(charts[id])charts[id].destroy();charts[id]=new Chart($(id),cfg);}

const anchorMap={};D.anchors.forEach(a=>{(anchorMap[a.q]=anchorMap[a.q]||[]).push(a.v);});
let mainRange='2018Q1';
function drawMain(C){
  const rows=T.filter(r=>r.q>=mainRange&&r.yoy!=null); const L=rows.map(r=>r.q);
  const fFrom=L.findIndex(q=>byQ[q].status==='forecast');
  const ds=[...band(rows.map(r=>r.hi),rows.map(r=>r.lo),C.orange),
    line('A-ATRI',rows.map(r=>r.yoy),C.orange,{borderWidth:2.6}),
    line('CPI rent (quarterly avg)',rows.map(r=>r.cpi),C.slate,{borderDash:[5,4],borderWidth:1.6})];
  if(D.full_bls) ds.push(line('BLS R-CPI-ATR',L.map(q=>D.full_bls[q]??null),C.ink,{borderWidth:1.2}));
  ds.push(pts('BLS published',L.map(q=>anchorMap[q]?anchorMap[q][0]:null),C));
  make('c-main',{type:'line',data:{labels:L,datasets:ds},options:baseOpts(C,L,{shadeFrom:fFrom,shadeLabel:'Forecast'})});
  legend('lg-main',[['A-ATRI',C.orange],['90% band',rgba(C.orange,0.25),'band'],['CPI rent',C.slate,'dash']].concat(D.full_bls?[['BLS R-CPI-ATR',C.ink]]:[]).concat([['BLS published print',C.ink,'pt']]));
}
function drawEdge(C){
  const rows=T.filter(r=>r.q>='2023Q1'); const L=rows.map(r=>r.q);
  const fFrom=L.findIndex(q=>byQ[q].status==='forecast');
  make('c-edge',{type:'line',data:{labels:L,datasets:[...band(rows.map(r=>r.hi),rows.map(r=>r.lo),C.orange),
    line('A-ATRI (combined)',rows.map(r=>r.yoy),C.orange,{borderWidth:2.6,pointRadius:2.5}),
    line('CPI inversion only',rows.map(r=>r.cpionly),C.orangeLt,{borderDash:[6,4],borderWidth:1.8}),
    line('Lease-roll only',rows.map(r=>r.leaseroll),C.teal,{borderDash:[2,3],borderWidth:2}),
    pts('BLS published',L.map(q=>anchorMap[q]?anchorMap[q][0]:null),C)]},
    options:baseOpts(C,L,{step:1,shadeFrom:fFrom,shadeLabel:'Forecast'})});
  legend('lg-edge',[['A-ATRI combined',C.orange],['90% band (backtest-scaled)',rgba(C.orange,0.25),'band'],['CPI inversion only',C.orangeLt,'dash'],['Lease-roll only (ZORI)',C.teal,'dot'],['BLS published',C.ink,'pt']]);
}
function shiftMap(map,L,k){return L.map((q,i)=>i-k>=0?(map[L[i-k]]??null):null);}
function qrange(a,b){const out=[];let y=+a.slice(0,4),q=+a.slice(-1);const by=+b.slice(0,4),bq=+b.slice(-1);while(y<by||(y===by&&q<=bq)){out.push(y+'Q'+q);q++;if(q>4){q=1;y++;}}return out;}
function drawLeads(C){
  const last=T[T.length-1].q; const L=qrange('2015Q2',last); const s=L.map(q=>byQ[q]?byQ[q].yoy:null);
  const fFrom=L.findIndex(q=>byQ[q]&&byQ[q].status==='forecast');
  const opts=baseOpts(C,L,{step:2,shadeFrom:fFrom});
  make('c-z',{type:'line',data:{labels:L,datasets:[line('A-ATRI',s,C.orange,{borderWidth:2.4}),line('ZORI y/y, +3Q',shiftMap(D.zori_q,L,3),C.teal,{borderWidth:1.8})]},options:opts});
  make('c-n',{type:'line',data:{labels:L,datasets:[line('A-ATRI',s,C.orange,{borderWidth:2.4}),line('BLS NTR y/y, +3Q',shiftMap(D.ntr_q,L,3),C.plum,{borderDash:[5,3],borderWidth:1.6,pointRadius:2.2,pointStyle:'rect'})]},options:baseOpts(C,L,{step:2,shadeFrom:fFrom})});
  legend('lg-z',[['A-ATRI',C.orange],['ZORI y/y, shifted +3Q',C.teal]]);
  legend('lg-n',[['A-ATRI',C.orange],['BLS NTR y/y, shifted +3Q',C.plum,'dash']]);
}
function drawW(C){
  const ws=D.wscan; const Lw=ws.map(d=>d.w.toFixed(2));
  make('c-wscan',{type:'bar',data:{labels:Lw,datasets:[{label:'Anchor RMSE',data:ws.map(d=>d.rmse),backgroundColor:ws.map(d=>Math.abs(d.w-P.w)<1e-9?C.orange:rgba(C.neutral,0.7)),borderRadius:{topLeft:3,topRight:3},borderSkipped:'bottom',barPercentage:0.86,categoryPercentage:0.9}]},
    options:{responsive:true,maintainAspectRatio:false,animation:false,plugins:{legend:{display:false},tooltip:{backgroundColor:C.card,titleColor:C.ink,bodyColor:C.ink2,borderColor:C.grid,borderWidth:1,titleFont:{family:'IBM Plex Mono'},bodyFont:{family:'IBM Plex Mono'},
      callbacks:{title:i=>'w = '+i[0].label,label:i=>{const d=ws[i.dataIndex];return [` RMSE ${f2(d.rmse)}pp`,` wedge ${sg(d.alpha)}pp/yr`,` peak ${d.peak} · trough ${d.trough}`];}}}},
      scales:{x:{grid:{display:false},ticks:{color:C.muted,font:{family:'IBM Plex Mono',size:10},callback(v){const l=Lw[v];return (+l*100)%25===0?l:'';},maxRotation:0,autoSkip:false},border:{color:C.grid}},y:{beginAtZero:true,grid:{color:C.grid},border:{display:false},ticks:{color:C.muted,font:{family:'IBM Plex Mono',size:11}}}}}});
  const keys=Object.keys(D.sens); const L=qrange('2019Q1',T[T.length-1].q);
  const ramp=[C.orangeLt,C.orange,C.orangeDk,C.orangeXdk];
  const ord=keys.map(k=>+k).sort((a,b)=>a-b).map(v=>v.toFixed(2));
  const ds=ord.map((k,i)=>line('w = '+k,L.map(q=>D.sens[k][q]??null),ramp[i],{borderWidth:Math.abs(+k-P.w)<1e-9?2.6:1.5}));
  ds.push(pts('BLS published',L.map(q=>anchorMap[q]?anchorMap[q][0]:null),C));
  make('c-sens',{type:'line',data:{labels:L,datasets:ds},options:baseOpts(C,L,{step:1,shadeFrom:L.findIndex(q=>byQ[q]&&byQ[q].status==='forecast')})});
  legend('lg-sens',ord.map((k,i)=>['w = '+k+(Math.abs(+k-P.w)<1e-9?' (chosen)':''),ramp[i]]));
}
function drawBT(C){
  const b=D.backtest; const L=b.map(d=>d.h===0?'Current qtr':'+'+d.h+'Q');
  const mk=(lab,key,col)=>({label:lab,data:b.map(d=>d[key]),backgroundColor:col,borderRadius:{topLeft:3,topRight:3},borderSkipped:'bottom',borderColor:C.card,borderWidth:{left:1,right:1},barPercentage:0.9,categoryPercentage:0.78});
  make('c-bt',{type:'bar',data:{labels:L,datasets:[mk('A-ATRI system','sys',C.orange),mk('CPI inversion only','cpionly',C.slate),mk('Naive: latest CPI y/y + wedge','naive',rgba(C.neutral,0.75))]},
    options:{responsive:true,maintainAspectRatio:false,animation:false,interaction:{mode:'index',intersect:false},plugins:{legend:{display:false},tooltip:{backgroundColor:C.card,titleColor:C.ink,bodyColor:C.ink2,borderColor:C.grid,borderWidth:1,titleFont:{family:'IBM Plex Mono'},bodyFont:{family:'IBM Plex Mono'},callbacks:{label:i=>` ${i.dataset.label}: ${f2(i.raw)}pp`}}},
      scales:{x:{grid:{display:false},ticks:{color:C.muted,font:{family:'IBM Plex Mono',size:11}},border:{color:C.grid}},y:{beginAtZero:true,grid:{color:C.grid},border:{display:false},ticks:{color:C.muted,font:{family:'IBM Plex Mono',size:11}}}}}});
  legend('lg-bt',[['A-ATRI system',C.orange],['CPI inversion only',C.slate],['Naive',C.neutral]]);
}
function drawSD(C){
  const rows=T.filter(r=>r.q>='2024Q3'); const L=rows.map(r=>r.q);
  make('c-sd',{type:'line',data:{labels:L,datasets:[line('Shutdown-aware (A-ATRI)',rows.map(r=>r.yoy),C.orange,{borderWidth:2.6,pointRadius:2.5}),line('Naive: published series at face value',rows.map(r=>r.naive_shutdown),C.slate,{borderDash:[5,4],borderWidth:1.8,pointRadius:2})]},
    options:baseOpts(C,L,{step:1,shadeFrom:L.findIndex(q=>byQ[q].status==='forecast'),shadeLabel:'Forecast'})});
  legend('lg-sd',[['Shutdown-aware (A-ATRI)',C.orange],['Naive: Oct carried forward, Apr as a normal month',C.slate,'dash']]);
}
function cpiHalf(d){const b=D.cpi_bt||[];let r=0;if(b.length){const hs=[{h:0,rmse:0}].concat(b);for(let i=1;i<hs.length;i++){if(d.h<=hs[i].h){const a=hs[i-1],c=hs[i];r=a.rmse+(c.rmse-a.rmse)*(d.h-a.h)/(c.h-a.h);break;}r=hs[hs.length-1].rmse;}}return 1.645*Math.max(d.sd,r);}
function drawCPI(C){
  const act=D.cpi_m, imp=D.icpi; const L=act.map(d=>d.m).concat(imp.map(d=>d.m));
  const A={};act.forEach(d=>A[d.m]=d.yoy); const I={};imp.forEach(d=>I[d.m]=d);
  const lastA=act[act.length-1];
  const at=(m,f)=>m===lastA.m?lastA.yoy:(I[m]?f(I[m]):null);
  const pub=L.map(m=>at(m,d=>d.yoy)), und=L.map(m=>at(m,d=>d.yoy_u));
  const hi=L.map(m=>at(m,d=>(d.yoy??d.yoy_u)+cpiHalf(d))), lo=L.map(m=>at(m,d=>(d.yoy??d.yoy_u)-cpiHalf(d)));
  make('c-cpi',{type:'line',data:{labels:L,datasets:[...band(hi,lo,C.slate,0.18),line('CPI rent, actual',L.map(m=>A[m]??null),C.slate,{borderWidth:2.4}),
    line('Implied, as BLS will print it',pub,C.slate,{borderDash:[5,4],borderWidth:2}),line('Implied, underlying',und,C.orange,{borderDash:[2,3],borderWidth:2})]},
    options:baseOpts(C,L,{step:1,shadeFrom:L.indexOf(imp.length?imp[0].m:null),shadeLabel:'Implied'})});
  legend('lg-cpi',[['CPI rent, actual',C.slate],['Implied, as BLS will print it',C.slate,'dash'],['Implied, underlying (base-corrected)',C.orange,'dot'],['90% band',rgba(C.slate,0.3),'band']]);
}
function drawAll(){const C=palette();drawMain(C);drawEdge(C);drawLeads(C);drawW(C);drawBT(C);drawSD(C);drawCPI(C);}

// ---------- text blocks & tables ----------
const nowQs=T.filter(r=>r.status==='nowcast').map(r=>r.q), fcQs=T.filter(r=>r.status==='forecast').map(r=>r.q);
$('edge-text').textContent=`Rent changes set in a quarter keep arriving in the CPI sample for up to six months, so the latest quarters (${nowQs.join(', ')}) are only partly revealed. A-ATRI combines the CPI inversion with a lease-roll model in which ZORI new-lease growth predicts the reset rate ${P.lr0.L} months later (slope ${f2(P.lr0.b)}, intercept ${f2(P.lr0.a)}pp). Forecast quarters (${fcQs.join(', ')}) rely on the lease-roll model alone.`;
$('edge-note').textContent=`Lease-roll errors are persistent (monthly AR(1) ρ = ${f2(P.rho)}), so the prior is correlated across months. Bands for nowcast and forecast quarters use the larger of the model's posterior uncertainty and the real-time backtest error at that horizon.`;
const xc=D.xcorr; const leads=Object.keys(xc['ZORI'].c).map(Number).sort((a,b)=>a-b);
$('t-xc').innerHTML=`<thead><tr><th>Lead (qtrs)</th>${leads.map(l=>`<th>${l>0?'+'+l:l}</th>`).join('')}<th>n</th></tr></thead><tbody>`+
  Object.entries(xc).map(([k,v])=>{const mx=Math.max(...leads.map(l=>v.c[l]??-9));return `<tr><td>${k}</td>${leads.map(l=>`<td class="${v.c[l]===mx?'hl':''}">${v.c[l]==null?'–':v.c[l].toFixed(2)}</td>`).join('')}<td>${v.n}</td></tr>`;}).join('')+'</tbody>';
const zb=Math.max(...leads.map(l=>xc['ZORI'].c[l]??-9)), zl=leads.find(l=>xc['ZORI'].c[l]===zb);
const nb=Math.max(...leads.map(l=>xc['BLS NTR'].c[l]??-9)), nl=leads.find(l=>xc['BLS NTR'].c[l]===nb);
$('zc').textContent=`peak r = ${zb.toFixed(2)} at +${zl}Q`; $('nc').textContent=`peak r = ${nb.toFixed(2)} at +${nl}Q`;
const V=D.validation, ch=D.checks;
$('val-text').textContent=`Calibration mode: ${D.sources.calib}. The wedge α is pinned by the two level prints; the re-dating weight w is picked on the four mature YoY prints; everything marked "check" is a test the model was not fitted to. Errors are A-ATRI minus published. The naive column is next quarter's CPI rent y/y plus the same wedge.`;
$('t-val').innerHTML=`<thead><tr><th>Quarter</th><th>Measure</th><th>Published</th><th>A-ATRI</th><th>Error</th><th>Naive</th><th>Naive err</th><th class="txt">Source</th><th class="txt">Role</th></tr></thead><tbody>`+
  V.map(r=>`<tr><td>${r.q}</td><td>${r.kind==='yoy'?'% y/y':'level'}</td><td>${r.kind==='yoy'?f2(r.pub):r.pub.toFixed(3)}</td><td>${r.kind==='yoy'?f2(r.syn):r.syn.toFixed(3)}</td><td>${sg(r.err)}</td><td>${f2(r.naive)}</td><td>${sg(r.nerr)}</td><td class="txt">${r.src}</td><td class="txt">${r.role}</td></tr>`).join('')+'</tbody>';
$('val-checks').innerHTML=`Out-of-sample checks: median y/y 2000Q4–2024Q4 published <span class="num">${f2(ch.median_pub)}</span>, A-ATRI <span class="num">${f2(ch.median_syn)}</span>. Series high published ${ch.peak_pub}, A-ATRI ${ch.peak_syn} (<span class="num">${f2(ch.peak_val)}</span>). Series low published ${ch.trough_pub}, A-ATRI ${ch.trough_syn} (<span class="num">${f2(ch.trough_val)}</span>). The 2023Q4 and 2024Q1 prints are first releases, which BLS warns are revised most; both A-ATRI and the naive benchmark sit below them.`;
const B=D.backtest;
$('bt-text').textContent=`At each quarter-end from Dec 2018, CPI and ZORI are cut off at that month and the full system is re-run (w and α held at final values). Each nowcast and forecast is compared with the final A-ATRI estimate for that quarter. The lease-roll layer earns its place beyond the current quarter. Raw model bands under-cover at longer horizons, so displayed bands use the larger of model uncertainty and backtest RMSE.`;
$('t-bt').innerHTML=`<thead><tr><th>Horizon</th><th>n</th><th>System</th><th>CPI only</th><th>Naive</th><th>Raw 90% cov.</th></tr></thead><tbody>`+B.map(d=>`<tr><td>${d.h===0?'Current':'+'+d.h+'Q'}</td><td>${d.n}</td><td>${f2(d.sys)}</td><td>${f2(d.cpionly)}</td><td>${f2(d.naive)}</td><td>${Math.round(100*d.raw_cov)}%</td></tr>`).join('')+'</tbody>';
const ic=D.icpi; const icl=ic.length?ic[ic.length-1]:null;
const cb=D.cpi_bt||[];
$('cpi-text').textContent=icl?`The A-ATRI reset rate, pushed through CPI's six-month panel arithmetic, implies CPI rent of primary residence at ${f2(icl.yoy)}% y/y by ${mlab(icl.m)}, from ${f2(H.cpi_now)}% in ${mlab(D.asof_cpi)}. Published twelve-month changes will not be clean in between: BLS's carried-forward October 2025 left published levels from October 2025 to March 2026 about ${f2(D.cpi_gap)}pp too low, so prints for November 2026 to March 2027 sit that much above the underlying rate${hump?` (peak ${f2(hump.yoy)}% in ${mlab(hump.m)})`:''}, and the effect unwinds in April 2027. October 2026 has no published base month.`+(cb.length?` In the real-time backtest the implied path missed actual CPI rent by ${cb.map(d=>f2(d.rmse)+'pp at +'+d.h+'m').join(', ')} (RMSE).`:''):'';
$('sources').innerHTML=[
  ['BLS: R-CPI-NTR and R-CPI-ATR research series (publication paused April 2026)','https://www.bls.gov/cpi/research-series/r-cpi-ntr.htm'],
  ['Adams, Loewenstein, Montag & Verbrugge, Disentangling Rent Index Differences (Cleveland Fed WP 22-38R)','https://www.clevelandfed.org/publications/working-paper/2023/wp-2238r-disentangling-rent-index-differences'],
  ['BLS: 2025 shutdown impact on the CPI, FAQ','https://www.bls.gov/cpi/additional-resources/2025-federal-government-shutdown-impact-cpi-faq.htm'],
  ['Monthly Labor Review: counterfactual imputation for October 2025 housing','https://www.bls.gov/opub/mlr/2026/article/counterfactual-imputation-approaches-for-the-housing-component-of-the-october-2025-cpi.htm'],
  ['CEIC: United States tenant rent index (published ATR anchors)','https://ceicdata.com/en/united-states/tenant-rent-index'],
  ['YCharts: BLS New Tenant Rent Index','https://ycharts.com/indicators/new_tenant_rent_index'],
  ['Zillow Research: ZORI','https://www.zillow.com/research/data/'],
  ['FRED: CPI rent of primary residence (CUUR0000SEHA)','https://fred.stlouisfed.org/series/CUUR0000SEHA']
].map(([t,u])=>`<li><a href="${u}" target="_blank" rel="noopener">${t}</a></li>`).join('')+`<li>Data used this build: CPI ${D.sources.cpi}; ZORI ${D.sources.zori}; NTR ${D.sources.ntr}.</li>`;
let dataAll=false;
function drawTable(){
  const rows=T.filter(r=>dataAll||r.q>='2018Q1').slice().reverse();
  $('t-data').innerHTML=`<thead><tr><th>Quarter</th><th class="txt" style="min-width:0">Status</th><th>Level</th><th>y/y %</th><th>90% band</th><th>q/q SAAR</th><th>CPI rent y/y</th><th>BLS published</th></tr></thead><tbody>`+
    rows.map(r=>`<tr><td>${r.q}</td><td class="txt" style="min-width:0"><span class="chip ${r.status}">${r.status}</span></td><td>${f2(r.level)}</td><td><b>${f2(r.yoy)}</b></td><td>${f2(r.lo)} – ${f2(r.hi)}</td><td>${f2(r.qoq)}</td><td>${f2(r.cpi)}</td><td>${r.pub==null?'':f2(r.pub)}</td></tr>`).join('')+'</tbody>';
}
function seg(a,b,on){$(a).setAttribute('aria-pressed',on?'true':'false');$(b).setAttribute('aria-pressed',on?'false':'true');}
$('r-full').onclick=()=>{mainRange='2000Q1';seg('r-full','r-recent',true);drawMain(palette());};
$('r-recent').onclick=()=>{mainRange='2018Q1';seg('r-full','r-recent',false);drawMain(palette());};
$('d-all').onclick=()=>{dataAll=true;seg('d-all','d-recent',true);drawTable();};
$('d-recent').onclick=()=>{dataAll=false;seg('d-all','d-recent',false);drawTable();};
$('copycsv').onclick=()=>{
  const cols=['q','status','level','yoy','lo','hi','qoq','cpi','cpionly','leaseroll','pub'];
  const csv=[cols.join(',')].concat(T.map(r=>cols.map(c=>r[c]==null?'':r[c]).join(','))).join('\n');
  const done=ok=>{$('copymsg').textContent=ok?`Copied ${T.length} rows to the clipboard.`:'Clipboard blocked here. Select the table and copy it manually.';};
  try{navigator.clipboard.writeText(csv).then(()=>done(true),()=>done(false));}catch(e){done(false);}
};
if(D.full_bls)$('lim-calib').textContent=`Calibration uses the full published BLS series (${D.sources.calib}); the last four quarters of the BLS vintage are excluded from fitting because BLS revises them most.`;
$('foot').textContent=`A-ATRI · generated ${D.generated} by rebuild.py · w=${P.w.toFixed(2)} · λ=${P.lam.toFixed(1)} · wedge ${sg(P.alpha)}pp/yr · Acheron Insights quantitative research. Not an official statistic.`;
drawTable(); drawAll();
const mq=window.matchMedia('(prefers-color-scheme: dark)');
if(mq.addEventListener)mq.addEventListener('change',drawAll);
new MutationObserver(drawAll).observe(document.documentElement,{attributes:true,attributeFilter:['data-theme']});
})();
</script>
"""

if __name__ == "__main__":
    main()
